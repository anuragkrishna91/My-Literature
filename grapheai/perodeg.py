"""
PeroDeg - by GrapheAI. Degradation analysis for perovskite cells & modules.

Round 1: universal data importer + long-term analysis (normalized PCE,
performance ratio, bilinear burn-in fit / PLR, T80) + the diurnal engine
(DPD / DPR per Paraskeva, Norton, Livera, ..., Krishna et al., ACS Energy
Lett. 2024, 9, 5081) + XGBoost power forecast + Claude degradation report
(optionally grounded in the paper corpus).

Run with:
    streamlit run perodeg.py --server.port 8504

Data lives in answers/perodeg_data/ (canonical CSVs) with import profiles
in answers/perodeg_profiles.json - both covered by the nightly backup.
"""

import datetime
import re
from pathlib import Path

import streamlit as st

ANSWERS_DIR = Path("answers")
PD_DATA_DIR = ANSWERS_DIR / "perodeg_data"
PD_PROFILES = ANSWERS_DIR / "perodeg_profiles.json"

MODELS = {
    "Fast (Haiku 4.5)": "claude-haiku-4-5",
    "Balanced (Sonnet 5)": "claude-sonnet-5",
    "Deep (Opus 5)": "claude-opus-5",
    "Frontier (Fable 5.1)": "claude-fable-5-1",
}

PRICES = {
    "claude-haiku-4-5": (1.00, 5.00),
    "claude-sonnet-4-6": (3.00, 15.00),
    "claude-sonnet-5": (2.00, 10.00),
    "claude-opus-4-8": (5.00, 25.00),
    "claude-opus-5": (5.00, 25.00),
    "claude-fable-5": (10.00, 50.00),
    "claude-fable-5-1": (10.00, 50.00),
}

SPEND_FILE = ANSWERS_DIR / "spend.json"


def _load_spend():
    import json
    if SPEND_FILE.exists():
        try:
            return json.loads(SPEND_FILE.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def _write_spend(data):
    import json
    try:
        ANSWERS_DIR.mkdir(exist_ok=True)
        SPEND_FILE.write_text(json.dumps(data), encoding="utf-8")
    except Exception:
        pass


def _add_spend(amount):
    if amount <= 0:
        return
    data = _load_spend()
    key = datetime.date.today().strftime("%Y-%m")
    data[key] = round(float(data.get(key, 0.0)) + amount, 5)
    _write_spend(data)


def _track_usage(n_in, n_out, model=None):
    u = st.session_state.setdefault("usage",
                                    {"in": 0, "out": 0, "calls": 0,
                                     "cost": 0.0})
    u["in"] += n_in
    u["out"] += n_out
    u["calls"] += 1
    p_in, p_out = PRICES.get(model, (0.0, 0.0))
    delta = n_in / 1e6 * p_in + n_out / 1e6 * p_out
    u["cost"] = u.get("cost", 0.0) + delta
    if st.session_state.get("backend") != "max":
        _add_spend(delta)


def _claude_cli_candidates():
    """Claude Code binaries the Agent SDK could run: the copy bundled inside
    claude-agent-sdk (which the SDK prefers by default) and any installed
    `claude` on PATH or in the usual locations."""
    import os
    import shutil
    from pathlib import Path
    cands = []
    try:
        import claude_agent_sdk
        b = (Path(claude_agent_sdk.__file__).parent / "_bundled"
             / ("claude.exe" if os.name == "nt" else "claude"))
        if b.is_file():
            cands.append(("bundled in claude-agent-sdk", str(b)))
    except Exception:
        pass
    seen = {c[1] for c in cands}
    home = Path.home()
    for p in [shutil.which("claude"), home / ".claude/local/claude",
              home / ".npm-global/bin/claude", "/usr/local/bin/claude",
              "/opt/homebrew/bin/claude", home / ".local/bin/claude"]:
        p = str(p) if p else ""
        if p and Path(p).is_file() and p not in seen:
            cands.append(("installed claude", p))
            seen.add(p)
    return cands


def _claude_cli_version(path):
    import re
    import subprocess
    try:
        out = subprocess.run([path, "--version"], capture_output=True, text=True,
                             timeout=20).stdout
        m = re.search(r"(\d+)\.(\d+)\.(\d+)", out or "")
        return m.group(0) if m else ""
    except Exception:
        return ""


def _ver_tuple(v):
    import re
    return tuple(int(x) for x in re.findall(r"\d+", str(v))[:3]) or (0,)


MAX_CLI_MIN_VERSION = "2.1.251"      # Fable 5.1 needs at least this Claude Code


def claude_code_status(refresh=False):
    """The newest Claude Code found (path, version, origin), cached for the
    session. Max mode runs through this binary, so a stale copy bundled in
    an old claude-agent-sdk no longer blocks new models when a newer
    `claude` is installed - and vice versa."""
    key = "_claude_code_status"
    if not refresh and key in st.session_state:
        return st.session_state[key]
    best = {}
    for origin, p in _claude_cli_candidates():
        v = _claude_cli_version(p)
        if v and (not best or _ver_tuple(v) > _ver_tuple(best["version"])):
            best = {"path": p, "version": v, "origin": origin}
    st.session_state[key] = best
    return best


def _max_update_command():
    import sys
    return f"{sys.executable} -m pip install -U claude-agent-sdk"


def _max_unsupported_hint(err, model):
    """Turn Claude Code's 'does not support this model; version X or newer
    is required' into a message that says exactly what to run."""
    import re
    m = re.search(r"does not support this model;?\s*version\s*([\d.]+)\s*or newer",
                  str(err), re.I)
    if not m:
        return None
    status = claude_code_status()
    have = status.get("version", "unknown")
    return (f"Claude Code {have} ({status.get('origin', 'not found')}) is too old "
            f"for {model}: version {m.group(1)} or newer is required. Fix in "
            f"Terminal, then restart the app:\n    {_max_update_command()}\n"
            "(claude-agent-sdk ships its own Claude Code; that command updates "
            "it. `claude update` works too - GrapheAI uses the newest Claude "
            "Code it finds.) Until then choose 'Fast (Haiku 4.5)' or switch the "
            "sidebar to 'API key'.")


def call_claude_max(system, user_msg, model, max_tokens=None):
    """Run one Claude call through the Claude Agent SDK (Claude Code), so it
    counts against the user's Claude Max subscription instead of API billing.

    Needs one-time setup (see MAX_SETUP.txt): `pip install claude-agent-sdk`
    (it bundles Claude Code) and a Claude Code login with the Max account.
    Personal use of your own subscription only. max_tokens is not enforced on
    this path (the CLI manages output length itself). New models need a
    recent Claude Code: the app runs the newest copy it can find.
    """
    try:
        from claude_agent_sdk import (query, ClaudeAgentOptions,
                                      AssistantMessage, TextBlock,
                                      ResultMessage)
    except ImportError as exc:
        raise RuntimeError(
            "Claude Max mode needs the Agent SDK. One-time setup (Terminal):\n"
            f"1) {_max_update_command()}\n"
            "2) Run `claude` once and log in with your Claude (Max) account\n"
            "Or switch the sidebar back to 'API key'. See MAX_SETUP.txt."
        ) from exc
    import asyncio
    status = claude_code_status()
    extra = {"cli_path": status["path"]} if status.get("path") else {}

    async def _run():
        opts = ClaudeAgentOptions(
            system_prompt=system,
            model=model,
            max_turns=1,
            # Pure text generation - keep Claude Code's tools out of the way.
            disallowed_tools=["Bash", "Edit", "Write", "Read", "Glob", "Grep",
                              "WebSearch", "WebFetch", "Task", "NotebookEdit"],
            **extra,
        )
        parts = []
        usage = None
        async for message in query(prompt=user_msg, options=opts):
            if isinstance(message, AssistantMessage):
                for block in message.content:
                    if isinstance(block, TextBlock):
                        parts.append(block.text)
            elif isinstance(message, ResultMessage):
                usage = getattr(message, "usage", None)
                if getattr(message, "is_error", False):
                    raise RuntimeError(getattr(message, "result", None)
                                       or "Claude Code returned an error.")
        return "".join(parts), usage

    try:
        text, usage = asyncio.run(_run())
    except Exception as exc:
        hint = _max_unsupported_hint(exc, model)
        if hint:
            raise RuntimeError(hint) from exc
        raise
    if not text.strip():
        raise RuntimeError(
            "Claude Code returned no text. If this mentions usage limits, "
            "your Max plan may be at its cap - switch the sidebar to "
            "'API key' to continue, or wait for the limit to reset.")
    try:
        _track_usage(int(usage.get("input_tokens", 0)),
                     int(usage.get("output_tokens", 0)))
    except Exception:
        _track_usage(0, 0)
    return text


THINKING_MODELS = ("claude-fable-5", "claude-mythos-5", "claude-opus-5",
                   "claude-opus-4-8", "claude-opus-4-7", "claude-sonnet-5")


def _effective_max_tokens(model, requested):
    mt = requested or 1500
    if any(str(model).startswith(m) for m in THINKING_MODELS):
        mt = max(mt, 32000)
    return mt


def _response_text(resp):
    if getattr(resp, "stop_reason", None) == "refusal":
        raise RuntimeError("Claude declined this request.")
    parts = [b.text for b in resp.content
             if getattr(b, "type", "") == "text" and getattr(b, "text", "")]
    if not parts:
        raise RuntimeError("The model returned no text.")
    return "".join(parts)


EFFORT_LEVELS = ("low", "medium", "high", "xhigh", "max")
STREAM_ABOVE = 8000      # stream responses above this many output tokens


def _model_family(model):
    m = str(model)
    if m.startswith(("claude-fable-5", "claude-mythos-5")):
        return "fable"          # thinking always on; depth via effort
    if m.startswith("claude-opus-5"):
        return "opus5"          # adaptive by default; effort accepted
    if m.startswith(("claude-opus-4-8", "claude-opus-4-7", "claude-sonnet-5")):
        return "adaptive"       # must ask for adaptive thinking
    if m.startswith(("claude-opus-4-6", "claude-sonnet-4-6")):
        return "adaptive46"     # adaptive; effort up to 'high'/'max'
    return "classic"            # Haiku 4.5 and older: no thinking


def _api_kwargs(model, system, user_msg, max_tokens):
    """Request parameters tuned per model family. Fable 5.1: thinking is
    always on and its depth is set with output_config.effort (the
    sidebar's 'Reasoning effort'); Opus 5 the same; Opus 4.8 / Sonnet 5
    need adaptive thinking requested explicitly; Haiku gets none."""
    fam = _model_family(model)
    kw = {"model": model,
          "max_tokens": _effective_max_tokens(model, max_tokens),
          "system": system,
          "messages": [{"role": "user", "content": user_msg}]}
    effort = st.session_state.get("effort", "xhigh")
    if effort not in EFFORT_LEVELS:
        effort = "xhigh"
    if fam in ("fable", "opus5"):
        kw["output_config"] = {"effort": effort}
    elif fam == "adaptive":
        kw["thinking"] = {"type": "adaptive"}
        kw["output_config"] = {"effort": effort}
    elif fam == "adaptive46":
        kw["thinking"] = {"type": "adaptive"}
        kw["output_config"] = {"effort": "high" if effort == "xhigh"
                               else effort}
    return kw


def _api_create(client, kw):
    """Send the request with graceful degradation: refusal fallbacks
    (beta), then effort/adaptive thinking, then the bare request - so an
    older installed SDK or a model that rejects a parameter still works.
    Long outputs are streamed to avoid HTTP timeouts."""
    import anthropic
    stream = kw["max_tokens"] > STREAM_ABOVE
    fam = _model_family(kw["model"])
    attempts = []
    if fam in ("fable", "opus5") and not st.session_state.get("_no_fallbacks"):
        attempts.append(("beta", dict(kw, betas=["server-side-fallback-2026-07-01"],
                                      fallbacks="default")))
    attempts.append(("plain", dict(kw)))
    attempts.append(("minimal", {k: kw[k] for k in
                                 ("model", "max_tokens", "system", "messages")}))
    last = None
    for kind, params in attempts:
        try:
            api = client.beta.messages if kind == "beta" else client.messages
            if stream:
                with api.stream(**params) as s:
                    return s.get_final_message()
            return api.create(**params)
        except TypeError as e:                 # installed SDK lacks a parameter
            last = e
            if kind == "beta":
                st.session_state["_no_fallbacks"] = True
            continue
        except anthropic.BadRequestError as e:
            last = e
            low = str(e).lower()
            if kind == "beta":
                st.session_state["_no_fallbacks"] = True
                continue
            if kind == "plain" and any(w in low for w in
                                       ("output_config", "effort", "thinking",
                                        "fallback", "beta", "unexpected")):
                continue
            raise
    raise last


def call_claude(api_key, system, user_msg, model, max_tokens=None):
    """Single Claude call via the selected backend; tracks token usage.
    Backend 'api': the Anthropic API (Fable 5.1 with effort-controlled
    reasoning, streaming, refusal fallbacks). Backend 'max': the Claude
    Agent SDK, billed to the Claude Max plan."""
    if st.session_state.get("backend") == "max":
        return call_claude_max(system, user_msg, model, max_tokens)
    import anthropic
    client = anthropic.Anthropic(api_key=api_key)
    resp = _api_create(client, _api_kwargs(model, system, user_msg,
                                           max_tokens))
    _track_usage(resp.usage.input_tokens, resp.usage.output_tokens, model)
    return _response_text(resp)


# --------------------------------------------------------------------------
# Word export (markdown -> real formatting)
# --------------------------------------------------------------------------
def _md_runs(par, text):
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`", text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        if m.group(1) is not None:
            par.add_run(m.group(1)).bold = True
        elif m.group(2) is not None:
            par.add_run(m.group(2)).italic = True
        else:
            r = par.add_run(m.group(3))
            r.font.name = "Consolas"
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def md_to_docx(doc, md):
    from docx.shared import Pt
    lines = md.splitlines()
    i, in_code = 0, False
    while i < len(lines):
        stripped = lines[i].strip()
        if stripped.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            r = doc.add_paragraph().add_run(lines[i])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            i += 1
            continue
        if not stripped or re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            i += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            tbl = []
            while i < len(lines):
                s = lines[i].strip()
                if not (s.startswith("|") and s.endswith("|")):
                    break
                cells = [c.strip() for c in s.strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    tbl.append(cells)
                i += 1
            if tbl:
                ncols = max(len(r) for r in tbl)
                t = doc.add_table(rows=len(tbl), cols=ncols)
                try:
                    t.style = "Table Grid"
                except Exception:
                    pass
                for ri, row in enumerate(tbl):
                    for ci in range(ncols):
                        cpar = t.cell(ri, ci).paragraphs[0]
                        _md_runs(cpar, row[ci] if ci < len(row) else "")
                        if ri == 0:
                            for r_ in cpar.runs:
                                r_.bold = True
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if m:
            h = doc.add_heading("", level=min(len(m.group(1)), 4))
            _md_runs(h, m.group(2).strip())
            i += 1
            continue
        m = re.match(r"^[-*+]\s+(.*)", stripped)
        if m:
            _md_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue
        m = re.match(r"^\d+[.)]\s+(.*)", stripped)
        if m:
            _md_runs(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue
        _md_runs(doc.add_paragraph(), stripped)
        i += 1


def save_report(name, text):
    ANSWERS_DIR.mkdir(exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    safe = re.sub(r"[^A-Za-z0-9 _-]", "", name)[:60].strip() or "report"
    (ANSWERS_DIR / f"PeroDeg_{safe}_{stamp}.md").write_text(
        f"# PeroDeg: {name}\n\n{text}\n", encoding="utf-8")
    try:
        from docx import Document
        doc = Document()
        doc.add_heading(f"PeroDeg: {name}", level=1)
        md_to_docx(doc, text)
        doc.save(ANSWERS_DIR / f"PeroDeg_{safe}_{stamp}.docx")
    except Exception:
        pass


# --------------------------------------------------------------------------
# Analysis core (pure functions - tested offline)
# --------------------------------------------------------------------------
CANON_FIELDS = [
    ("timestamp", "Timestamp (required)"),
    ("device", "Device ID"),
    ("pce", "PCE (%)"),
    ("pmax", "Pmax / power (W)"),
    ("isc", "Isc (A or mA)"),
    ("voc", "Voc (V)"),
    ("ff", "Fill factor"),
    ("imp", "Imp (A or mA)"),
    ("vmp", "Vmp (V)"),
    ("gni", "Irradiance POA/GNI (W/m2)"),
    ("t_mod", "Module temperature (C)"),
    ("t_amb", "Ambient temperature (C)"),
    ("rh", "Relative humidity (%)"),
    ("sweep", "Sweep direction (fwd/rev)"),
    ("bias", "Bias hold (voc/mpp)"),
]
METRIC_FIELDS = ["pce", "pmax", "isc", "voc", "ff", "imp", "vmp"]


def normalize_df(raw, colmap, dayfirst=False):
    """raw dataframe + {canonical: source column} -> canonical dataframe,
    sorted by time, numeric metrics coerced."""
    import pandas as pd
    out = {}
    for canon, src in colmap.items():
        if src and src in raw.columns:
            out[canon] = raw[src]
    df = pd.DataFrame(out)
    if "timestamp" not in df.columns:
        raise ValueError("A timestamp column must be mapped.")
    df["timestamp"] = pd.to_datetime(df["timestamp"], errors="coerce",
                                     dayfirst=dayfirst)
    df = df.dropna(subset=["timestamp"])
    for f in METRIC_FIELDS + ["gni", "t_mod", "t_amb", "rh"]:
        if f in df.columns:
            df[f] = pd.to_numeric(df[f], errors="coerce")
    if "device" not in df.columns:
        df["device"] = "device1"
    df["device"] = df["device"].astype(str)
    for f in ("sweep", "bias"):
        if f in df.columns:
            df[f] = df[f].astype(str).str.lower().str.strip()
    df = df.sort_values("timestamp").reset_index(drop=True)
    df["date"] = df["timestamp"].dt.date
    return df


def daily_series(df, metric, norm_days=3):
    """Daily mean of a metric, plus a column normalized to the mean of the
    first `norm_days` days. Returns DataFrame(date, value, norm, day)."""
    import pandas as pd
    d = (df.dropna(subset=[metric]).groupby("date")[metric]
         .mean().reset_index().rename(columns={metric: "value"}))
    if d.empty:
        return d
    ref = d["value"].head(norm_days).mean()
    d["norm"] = d["value"] / ref if ref else None
    d0 = d["date"].iloc[0]
    d["day"] = d["date"].apply(lambda x: (x - d0).days)
    return d


def bilinear_fit(days, values):
    """Continuous two-segment least-squares fit:
    y = a + b1*min(x,k) + b2*max(x-k, 0), knee k searched over the data.
    Returns (k, a, b1, b2, fitted) or None if too little data."""
    import numpy as np
    x = np.asarray(days, dtype=float)
    y = np.asarray(values, dtype=float)
    ok = np.isfinite(x) & np.isfinite(y)
    x, y = x[ok], y[ok]
    if len(x) < 10:
        return None
    lo = np.quantile(x, 0.1)
    hi = np.quantile(x, 0.9)
    candidates = np.unique(x[(x >= lo) & (x <= hi)])
    if len(candidates) > 120:
        candidates = np.linspace(lo, hi, 120)
    best = None
    for k in candidates:
        A = np.column_stack([np.ones_like(x), np.minimum(x, k),
                             np.maximum(x - k, 0.0)])
        coef, *_ = np.linalg.lstsq(A, y, rcond=None)
        sse = float(((A @ coef - y) ** 2).sum())
        if best is None or sse < best[0]:
            best = (sse, k, coef)
    _, k, coef = best
    a, b1, b2 = (float(c) for c in coef)
    import numpy as np2
    A = np2.column_stack([np2.ones_like(x), np2.minimum(x, k),
                          np2.maximum(x - k, 0.0)])
    fitted = A @ np2.array([a, b1, b2])
    return float(k), a, b1, b2, fitted


def t80(daily, roll=7):
    """First day the rolling-median normalized value drops to <= 0.80.
    Returns day number or None."""
    d = daily.dropna(subset=["norm"]).copy()
    if d.empty:
        return None
    d["smooth"] = d["norm"].rolling(roll, min_periods=1,
                                    center=True).median()
    hit = d[d["smooth"] <= 0.80]
    return int(hit["day"].iloc[0]) if not hit.empty else None


def _fit_t80(k, a, b1, b2, x_max, target=0.80):
    """Day the fitted bilinear model crosses `target`, or None if it
    never does within [0, x_max]."""
    if b1 < 0:
        x = (target - a) / b1
        if 0 <= x <= k:
            return float(x)
    y_k = a + b1 * k
    if y_k > target and b2 < 0:
        x = k + (target - y_k) / b2
        if x <= x_max:
            return float(x)
    return None


def bootstrap_bilinear(days, values, n_boot=300, seed=0):
    """Pairs-bootstrap 95% confidence intervals for the bilinear PLR
    fit: burn-in slope, stable slope (both %/month), knee day, and the
    model-implied T80. Returns a dict or None if the base fit fails."""
    import numpy as np
    x = np.asarray(days, dtype=float)
    y = np.asarray(values, dtype=float)
    ok = np.isfinite(x) & np.isfinite(y)
    x, y = x[ok], y[ok]
    if bilinear_fit(x, y) is None:
        return None
    rng = np.random.default_rng(seed)
    b1s, b2s, knees, t80s = [], [], [], []
    x_max = float(x.max())
    for _ in range(int(n_boot)):
        idx = rng.integers(0, len(x), len(x))
        fit = bilinear_fit(x[idx], y[idx])
        if fit is None:
            continue
        k, a, b1, b2, _f = fit
        b1s.append(b1 * 3000.0)          # fraction/day -> %/month
        b2s.append(b2 * 3000.0)
        knees.append(k)
        t = _fit_t80(k, a, b1, b2, x_max)
        t80s.append(t if t is not None else np.inf)

    def _ci(vals):
        v = np.asarray(vals, dtype=float)
        return (float(np.percentile(v, 2.5)),
                float(np.percentile(v, 97.5)))

    t80_arr = np.asarray(t80s, dtype=float)
    frac_reached = float(np.mean(np.isfinite(t80_arr))) if len(
        t80_arr) else 0.0
    t80_ci = (_ci(t80_arr[np.isfinite(t80_arr)])
              if frac_reached >= 0.5 else None)
    return {"n_ok": len(b1s), "n_boot": int(n_boot),
            "burnin_plr_ci": _ci(b1s), "stable_plr_ci": _ci(b2s),
            "knee_ci": _ci(knees), "t80_ci": t80_ci,
            "t80_frac_reached": frac_reached}


def compute_dpd_dpr(df, metric="pce", g_lo=350, g_hi=450, dt_max=None,
                    min_gap_h=4):
    """Diurnal performance degradation / overnight recovery, after
    ACS Energy Lett. 2024, 9, 5081 (eqs 1-2):
      DPD(d)  = [M(d)  - E(d)] / M(day1)
      DPR(d)  = [M(d+1) - E(d)] / M(day1)
    where M/E are the first/last samples of each day inside the
    irradiance window [g_lo, g_hi]. Values returned in percent.
    dt_max: optionally require |T_morn - T_eve| <= dt_max (module temp,
    falling back to ambient)."""
    import pandas as pd
    need = {"timestamp", metric, "gni"}
    if not need <= set(df.columns):
        return pd.DataFrame()
    w = df.dropna(subset=[metric, "gni"])
    w = w[(w["gni"] >= g_lo) & (w["gni"] <= g_hi)]
    tcol = ("t_mod" if "t_mod" in df.columns
            else ("t_amb" if "t_amb" in df.columns else None))
    days = []
    for date, g in w.groupby("date"):
        g = g.sort_values("timestamp")
        first, last = g.iloc[0], g.iloc[-1]
        gap_h = (last["timestamp"] - first["timestamp"]).total_seconds() / 3600
        if gap_h < min_gap_h:
            continue
        if dt_max is not None and tcol is not None:
            tm, te = first.get(tcol), last.get(tcol)
            if pd.notna(tm) and pd.notna(te) and abs(tm - te) > dt_max:
                continue
        rec = {"date": date, "morning": float(first[metric]),
               "evening": float(last[metric])}
        if tcol is not None:
            tm, te = first.get(tcol), last.get(tcol)
            rec["t_day"] = (float((tm + te) / 2)
                            if pd.notna(tm) and pd.notna(te) else None)
        days.append(rec)
    if not days:
        return pd.DataFrame()
    dd = pd.DataFrame(days).sort_values("date").reset_index(drop=True)
    ref = dd["morning"].iloc[0]
    if not ref:
        return pd.DataFrame()
    dd["dpd"] = (dd["morning"] - dd["evening"]) / ref * 100
    nxt = dd["morning"].shift(-1)
    consec = (pd.to_datetime(dd["date"]).shift(-1)
              - pd.to_datetime(dd["date"])).dt.days == 1
    dd["dpr"] = ((nxt - dd["evening"]) / ref * 100).where(consec)
    dd["month"] = pd.to_datetime(dd["date"]).dt.strftime("%Y-%m")
    return dd


def perf_ratio_daily(df, pnom_w):
    """Daily performance ratio = (mean Pmax / Pnom) / (mean GNI / 1000),
    over daylight samples (GNI > 50)."""
    import pandas as pd
    if not {"pmax", "gni"} <= set(df.columns) or not pnom_w:
        return pd.DataFrame()
    w = df.dropna(subset=["pmax", "gni"])
    w = w[w["gni"] > 50]
    d = w.groupby("date").agg(p=("pmax", "mean"),
                              g=("gni", "mean")).reset_index()
    d["pr"] = (d["p"] / pnom_w) / (d["g"] / 1000.0)
    d0 = d["date"].iloc[0] if not d.empty else None
    if d0 is not None:
        d["day"] = d["date"].apply(lambda x: (x - d0).days)
    return d


def nrmse_nmbe(actual, pred):
    import numpy as np
    a = np.asarray(actual, float)
    p = np.asarray(pred, float)
    ok = np.isfinite(a) & np.isfinite(p)
    a, p = a[ok], p[ok]
    if len(a) == 0 or a.mean() == 0:
        return None, None
    rmse = float(np.sqrt(((p - a) ** 2).mean()))
    mbe = float((p - a).mean())
    return rmse / a.mean() * 100, mbe / a.mean() * 100


def extract_iv_params(v, i):
    """IV-curve parameters from one sweep. Handles either current sign
    convention (photocurrent is made positive). Returns dict with isc,
    voc, pmax, vmp, imp, ff - or None if the curve is unusable."""
    import numpy as np
    v = np.asarray(v, float)
    i = np.asarray(i, float)
    ok = np.isfinite(v) & np.isfinite(i)
    v, i = v[ok], i[ok]
    if len(v) < 8:
        return None
    order = np.argsort(v)
    v, i = v[order], i[order]
    # make photocurrent positive: current near V=0 should be > 0
    j0 = np.argmin(np.abs(v))
    if i[max(0, j0 - 2):j0 + 3].mean() < 0:
        i = -i
    # Isc: interpolate at V=0 (or nearest point if 0 not spanned)
    if v[0] <= 0 <= v[-1]:
        isc = float(np.interp(0.0, v, i))
    else:
        isc = float(i[j0])
    # Voc: first zero crossing of current for V > 0
    pos = v > 0
    vp, ip = v[pos], i[pos]
    cross = np.where(np.diff(np.sign(ip)) < 0)[0]
    if len(cross):
        k = cross[0]
        # linear interpolation between the two points around the crossing
        v1, v2, i1, i2 = vp[k], vp[k + 1], ip[k], ip[k + 1]
        voc = float(v1 + (v2 - v1) * (i1 / (i1 - i2)))
    elif (ip > 0).all():
        voc = float(vp[-1])
    else:
        return None
    # power in the operating quadrant
    quad = (v >= 0) & (v <= voc) & (i >= 0)
    if quad.sum() < 3:
        return None
    p = v[quad] * i[quad]
    kmax = int(np.argmax(p))
    pmax = float(p[kmax])
    vmp = float(v[quad][kmax])
    imp = float(i[quad][kmax])
    ff = pmax / (isc * voc) if isc > 0 and voc > 0 else None
    return {"isc": isc, "voc": voc, "pmax": pmax, "vmp": vmp,
            "imp": imp, "ff": (float(ff) if ff else None)}


def sweep_direction(v):
    """'fwd' if voltage predominantly increases over the sweep."""
    import numpy as np
    v = np.asarray(v, float)
    if len(v) < 2:
        return "fwd"
    return "fwd" if v[-1] >= v[0] else "rev"


def daily_energy(df, pnom_w=None):
    """Trapezoidal integration of pmax over each day's timestamps.
    Returns DataFrame(date, energy_wh, hours, ref_kwh_m2?, pr?)."""
    import numpy as np
    import pandas as pd
    if "pmax" not in df.columns:
        return pd.DataFrame()
    _trapz = np.trapezoid if hasattr(np, "trapezoid") else np.trapz
    out = []
    for date, g in df.dropna(subset=["pmax"]).groupby("date"):
        g = g.sort_values("timestamp")
        if len(g) < 3:
            continue
        t = ((g["timestamp"] - g["timestamp"].iloc[0])
             .dt.total_seconds().values / 3600.0)
        e_wh = float(_trapz(g["pmax"].values, t))
        rec = {"date": date, "energy_wh": e_wh,
               "hours": float(t[-1])}
        if "gni" in g.columns and g["gni"].notna().sum() >= 3:
            gg = g.dropna(subset=["gni"]).sort_values("timestamp")
            tg = ((gg["timestamp"] - gg["timestamp"].iloc[0])
                  .dt.total_seconds().values / 3600.0)
            ref = float(_trapz(gg["gni"].values, tg)) / 1000.0
            rec["ref_kwh_m2"] = ref
            if pnom_w and ref > 0:
                rec["pr"] = (e_wh / pnom_w) / ref
        out.append(rec)
    d = pd.DataFrame(out)
    if not d.empty:
        d["month"] = pd.to_datetime(d["date"]).dt.strftime("%Y-%m")
    return d


def day_night_analysis(dd):
    """Reversible/irreversible decomposition from a DPD/DPR table
    (output of compute_dpd_dpr): irreversible rate = linear trend of the
    normalized morning values; reversible amplitude = median DPD;
    recovery efficiency = median DPR / median DPD. Also detects data
    gaps > 2 days and quantifies recovery across them (dark-storage
    effect). Returns (summary dict, gaps DataFrame)."""
    import numpy as np
    import pandas as pd
    if dd is None or len(dd) < 5:
        return None, pd.DataFrame()
    d = dd.sort_values("date").reset_index(drop=True)
    ref = d["morning"].iloc[0]
    if not ref:
        return None, pd.DataFrame()
    dates = pd.to_datetime(d["date"])
    day_idx = (dates - dates.iloc[0]).dt.days.values.astype(float)
    mnorm = d["morning"].values / ref
    slope_per_day = float(np.polyfit(day_idx, mnorm, 1)[0])
    summary = {
        "days_span": int(day_idx[-1]),
        "irreversible_rate_pct_per_month": round(slope_per_day
                                                 * 30 * 100, 2),
        "reversible_amplitude_pct": round(float(d["dpd"].median()), 2),
        "recovery_efficiency_pct": (
            round(float(d["dpr"].median() / d["dpd"].median() * 100), 1)
            if d["dpd"].median() else None),
        "total_morning_loss_pct": round((1 - mnorm[-1]) * 100, 1),
    }
    gaps = []
    for k in range(len(d) - 1):
        gap_days = (dates.iloc[k + 1] - dates.iloc[k]).days
        if gap_days > 2:
            before = d["evening"].iloc[k]
            after = d["morning"].iloc[k + 1]
            gaps.append({
                "gap_start": d["date"].iloc[k],
                "gap_end": d["date"].iloc[k + 1],
                "days": gap_days,
                "before (evening)": round(float(before), 3),
                "after (morning)": round(float(after), 3),
                "recovery_pct_of_initial": round(
                    float((after - before) / ref * 100), 2),
            })
    return summary, pd.DataFrame(gaps)


def load_image_gray(data, max_side=800):
    """Image bytes -> downsampled float grayscale numpy array.
    Handles 8/16-bit PNG/TIFF/JPG."""
    import io as _io
    import numpy as np
    from PIL import Image
    img = Image.open(_io.BytesIO(data))
    if img.mode not in ("L", "I", "I;16", "F"):
        img = img.convert("L")
    arr = np.asarray(img, dtype=float)
    if arr.ndim == 3:
        arr = arr.mean(axis=2)
    k = max(1, int(max(arr.shape) / max_side))
    return arr[::k, ::k]


def image_stats(arr, ref_median=None, dark_ratio=0.5, hot_sigma=3.0):
    """Per-image degradation metrics.
    dark_frac: fraction of pixels below dark_ratio x the REFERENCE
    image's median (PL/EL dark-area growth).
    hot_frac: fraction above mean + hot_sigma*sigma of THIS image
    (DLIT hotspots)."""
    import numpy as np
    a = np.asarray(arr, float)
    out = {"mean": float(a.mean()), "median": float(np.median(a)),
           "std": float(a.std())}
    if ref_median:
        out["dark_frac"] = float((a < dark_ratio * ref_median).mean())
    hot_thr = a.mean() + hot_sigma * a.std()
    out["hot_frac"] = float((a > hot_thr).mean())
    out["hot_intensity_ratio"] = (float(a[a > hot_thr].mean()
                                        / a.mean())
                                  if (a > hot_thr).any() and a.mean()
                                  else None)
    return out


KB_EV = 8.617333e-5  # Boltzmann constant, eV/K


def arrhenius_fit(temps_c, rates):
    """ln k = ln A - Ea/(kB*T). rates must be positive degradation-rate
    magnitudes (same units throughout, e.g. %/month).
    Returns (Ea_eV, lnA, r2) or None if fewer than 3 usable points."""
    import numpy as np
    T = np.asarray(temps_c, float) + 273.15
    k = np.asarray(rates, float)
    ok = np.isfinite(T) & np.isfinite(k) & (k > 0)
    if ok.sum() < 3:
        return None
    x = 1.0 / T[ok]
    y = np.log(k[ok])
    b, a = np.polyfit(x, y, 1)
    Ea = -b * KB_EV
    yhat = b * x + a
    ss_res = float(((y - yhat) ** 2).sum())
    ss_tot = float(((y - y.mean()) ** 2).sum())
    r2 = 1 - ss_res / ss_tot if ss_tot else 1.0
    return float(Ea), float(a), float(r2)


def acceleration_factor(ea_ev, t1_c, t2_c):
    """k(T2)/k(T1) for activation energy Ea."""
    import numpy as np
    return float(np.exp(-ea_ev / KB_EV
                        * (1 / (t2_c + 273.15) - 1 / (t1_c + 273.15))))


def isos_table(meta, metrics):
    """Deterministic ISOS-style summary table (markdown) from the
    session's computed metrics + user-entered conditions. Follows the
    reporting spirit of the ISOS consensus (Khenkin et al., Nat. Energy
    2020): every figure of merit next to its conditions."""
    lt = metrics.get("long_term", {})
    di = metrics.get("diurnal", {})
    rv = metrics.get("reversibility", {})
    rows = [
        ("ISOS protocol", meta.get("protocol", "-")),
        ("Sample description", meta.get("sample", "-")),
        ("Location / setup", meta.get("location", "-")),
        ("Bias between measurements", meta.get("bias", "-")),
        ("Initial performance", meta.get("initial", "-")),
        ("Exposure duration", f"{lt.get('days', '?')} days"),
        ("Final performance (normalized)", lt.get("final_norm", "-")),
        ("T80", (f"day {lt['t80_day']}"
                 if lt.get("t80_day") is not None else "not reached")),
        ("PLR, burn-in phase",
         (f"{lt['burnin_plr_pct_per_month']} %/month"
          if lt.get("burnin_plr_pct_per_month") is not None else "-")),
        ("PLR, stable phase",
         (f"{lt['stable_plr_pct_per_month']} %/month"
          if lt.get("stable_plr_pct_per_month") is not None else "-")),
        ("Burn-in knee", (f"day {lt['knee_day']:.0f}"
                          if lt.get("knee_day") is not None else "-")),
        ("Median diurnal degradation (DPD)",
         (f"{di['median_dpd_pct']}% (window {di.get('window_wm2', '?')} "
          f"W/m², {di.get('valid_days', '?')} days)"
          if di.get("median_dpd_pct") is not None else "-")),
        ("Median overnight recovery (DPR)",
         (f"{di['median_dpr_pct']}%"
          if di.get("median_dpr_pct") is not None else "-")),
        ("Irreversible rate (morning trend)",
         (f"{rv['irreversible_rate_pct_per_month']} %/month"
          if rv.get("irreversible_rate_pct_per_month") is not None
          else "-")),
        ("Recovery efficiency",
         (f"{rv['recovery_efficiency_pct']}%"
          if rv.get("recovery_efficiency_pct") is not None else "-")),
    ]
    lines = ["| Quantity | Value |", "|---|---|"]
    lines += [f"| {q} | {v} |" for q, v in rows]
    lines.append("")
    lines.append("*Metastability note (ISOS consensus): perovskite "
                 "figures of merit depend on measurement conditions and "
                 "prior bias/light history; the conditions above are "
                 "part of the result.*")
    return "\n".join(lines)


def save_pub_fig(fig, name):
    """Save a matplotlib figure as 300-dpi PNG + SVG into
    answers/figures_out/. Returns the PNG path."""
    out = ANSWERS_DIR / "figures_out"
    out.mkdir(parents=True, exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    safe = re.sub(r"[^A-Za-z0-9_-]", "_", name)[:60]
    png = out / f"{safe}_{stamp}.png"
    fig.savefig(png, dpi=300, bbox_inches="tight",
                facecolor="white")
    fig.savefig(out / f"{safe}_{stamp}.svg", bbox_inches="tight",
                facecolor="white")
    return png


def _pub_ax(ax):
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(labelsize=9)
    return ax


# Editable monthly plane-of-array insolation presets (kWh/m2/month).
YIELD_PRESETS = {
    "Nicosia, CY (sunny)": [95, 105, 155, 185, 215, 230, 240, 220, 180,
                            140, 100, 85],
    "Genk, BE (temperate)": [25, 45, 90, 130, 160, 165, 160, 140, 100,
                             60, 30, 20],
    "Custom": [100] * 12,
}
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug",
          "Sep", "Oct", "Nov", "Dec"]


# --------------------------------------------------------------------------
# Storage: datasets + import profiles
# --------------------------------------------------------------------------
def load_profiles():
    import json
    if PD_PROFILES.exists():
        try:
            return json.loads(PD_PROFILES.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


def save_profiles(p):
    import json
    try:
        ANSWERS_DIR.mkdir(exist_ok=True)
        PD_PROFILES.write_text(json.dumps(p, ensure_ascii=False),
                               encoding="utf-8")
    except Exception:
        pass


def list_datasets():
    if not PD_DATA_DIR.exists():
        return []
    return sorted(p.stem for p in PD_DATA_DIR.glob("*.csv"))


def load_dataset(name):
    import pandas as pd
    df = pd.read_csv(PD_DATA_DIR / f"{name}.csv",
                     parse_dates=["timestamp"])
    df["date"] = df["timestamp"].dt.date
    return df


DEG_REPORT_SYSTEM = """\
You write the degradation-analysis section for a perovskite PV
device/module, from computed metrics (and optionally literature excerpts
from the author's own paper library).

Structure (markdown):
1. **Performance summary** - initial performance, loss over the period,
   T80 status, the bilinear fit: burn-in rate, stable-phase rate, knee.
2. **Diurnal behaviour** - DPD/DPR statistics, their ratio/correlation,
   temperature dependence, and which parameter (current vs voltage vs
   FF) drives the diurnal changes if the data shows it.
3. **Mechanistic interpretation** - hypotheses CONSISTENT with these
   signatures (e.g. reversible ion-migration behaviour vs irreversible
   chemical/electrode degradation), each clearly labelled a hypothesis.
   Where literature excerpts are provided, cite them as [n]; without
   excerpts, keep interpretation qualitative and flag it as ungrounded.
4. **Recommendations** - measurements or protocol changes that would
   discriminate between the hypotheses (e.g. bias-hold variation,
   temperature-controlled indoor cycling, EL/PL imaging).

Rules: use ONLY the provided numbers - never invent values; distinguish
reversible from irreversible losses explicitly; hypotheses are for the
researcher's judgement, not conclusions; keep ISOS-style reporting
discipline (state conditions alongside every number)."""


# --------------------------------------------------------------------------
# Page, theme, sidebar
# --------------------------------------------------------------------------
st.set_page_config(page_title="PeroDeg - by GrapheAI",
                   page_icon="🔋", layout="wide")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Serif:wght@600&family=IBM+Plex+Sans:wght@400;500;600&family=IBM+Plex+Mono:wght@400;500&display=swap');
html, body, [class*="css"], .stMarkdown, p, li {
    font-family: 'IBM Plex Sans', 'Segoe UI', sans-serif;
    font-size: 16.5px; color: #E6EAF0;
}
.stApp { background: #12161C; }
.pdg-header { border-bottom: 3px solid #FF6B3D; padding-bottom: 12px;
              margin-bottom: 10px; }
.pdg-title  { font-family: 'IBM Plex Serif', Georgia, serif;
              font-size: 2.6rem; font-weight: 600; color: #F4F6F9;
              margin: 0; letter-spacing: -0.5px; }
.pdg-sub    { font-size: 0.82rem; color: #8E99A8; margin-top: 6px;
              text-transform: uppercase; letter-spacing: 2px; }
.pdg-sub b  { color: #FF8A5C; font-weight: 600; }
.stTabs [data-baseweb="tab-list"] { border-bottom: 1px solid #242D38; }
.stTabs [data-baseweb="tab"] { font-size: 1.05rem; font-weight: 500;
    padding: 12px 18px; color: #8E99A8; }
.stTabs [data-baseweb="tab"]:hover { color: #FFB08F; }
.stTabs [aria-selected="true"] { color: #FF8A5C; font-weight: 600; }
.stTabs [data-baseweb="tab-highlight"] { background-color: #FF6B3D;
                                         height: 3px; }
section[data-testid="stSidebar"] { background: #171D25;
    border-right: 1px solid #242D38; }
section[data-testid="stSidebar"] h1 { font-family: 'IBM Plex Serif', serif;
    font-size: 1.4rem; color: #F4F6F9; }
h2 { font-family: 'IBM Plex Serif', Georgia, serif; font-weight: 600;
     color: #DFE5EC; }
h3 { color: #C3CBD6; }
[data-testid="stForm"] { background: #1A212B; border: 1px solid #263140;
    border-radius: 14px; padding: 1.1rem 1.3rem .9rem; }
details { border: 1px solid #263140; border-radius: 12px;
          background: #1A212B; }
[data-testid="stDataFrame"] { border: 1px solid #263140;
                              border-radius: 12px; }
[data-testid="stMetric"] { background: #1A212B; border: 1px solid #263140;
    border-radius: 12px; padding: .7rem .95rem; }
[data-testid="stMetricValue"] { font-family: 'IBM Plex Mono', monospace;
                                color: #FF8A5C; }
[data-testid="stMetricLabel"] { color: #8E99A8; text-transform: uppercase;
    letter-spacing: 1px; font-size: .78rem; }
.stButton>button, .stDownloadButton>button, .stFormSubmitButton>button {
    font-size: 1.0rem; font-weight: 600; border-radius: 10px;
    padding: 0.5rem 1.25rem; }
</style>
""", unsafe_allow_html=True)

st.markdown(
    "<div class='pdg-header'>"
    "<p class='pdg-title'>🔋 PeroDeg</p>"
    "<p class='pdg-sub'>degradation analytics · by <b>GrapheAI</b> · "
    "developed by <b>Dr. Anurag Krishna</b></p>"
    "</div>",
    unsafe_allow_html=True)

with st.sidebar:
    st.title("🔋 PeroDeg")
    st.caption("by GrapheAI · Dr. Anurag Krishna")
    backend_label = st.radio(
        "Claude access", ["API key (pay per use)",
                          "Claude Max subscription (needs Claude Code)"])
    st.session_state["backend"] = ("max" if "Max" in backend_label
                                   else "api")
    if st.session_state["backend"] == "api":
        api_key = st.text_input("Anthropic API key", type="password")
    else:
        api_key = "claude-max-subscription"
        _cc = claude_code_status()
        if not _cc:
            st.warning("Claude Code not found for Max mode - run "
                       f"`{_max_update_command()}` and log in with `claude`.")
        elif _ver_tuple(_cc["version"]) < _ver_tuple(MAX_CLI_MIN_VERSION):
            st.warning(f"Claude Code {_cc['version']} ({_cc['origin']}) is older "
                       f"than {MAX_CLI_MIN_VERSION}, which Fable 5.1 needs. Run in "
                       f"Terminal, then restart: `{_max_update_command()}`. Until "
                       "then choose Haiku 4.5 or the API-key backend.")
        else:
            st.caption(f"Claude Code {_cc['version']} ({_cc['origin']})")
    st.session_state["effort"] = st.select_slider(
        "Reasoning effort", options=list(EFFORT_LEVELS),
        value=st.session_state.get("effort", "xhigh"),
        help="Thinking depth for Fable 5.1 / Opus 5 / Sonnet 5 on the API "
             "backend. xhigh = professor-grade default; max = when "
             "correctness matters more than time; low = quick lookups.")
    st.markdown("---")
    dsets = list_datasets()
    st.caption(f"Datasets stored: {len(dsets)}")
    u = st.session_state.get("usage")
    if u:
        st.caption(f"Session: {u['calls']} calls - "
                   f"{u['in']:,}/{u['out']:,} tokens")

try:
    import pandas as _pd
except Exception:
    st.error("pandas is required.")
    st.stop()
try:
    import altair as alt
except Exception:
    alt = None

(tab_data, tab_iv, tab_long, tab_diurnal, tab_rev, tab_cmp, tab_arr,
 tab_yield, tab_img, tab_forecast, tab_report) = st.tabs(
    ["📥 Data", "📈 I-V curves", "📉 Long-term", "🌗 Diurnal",
     "♻️ Reversibility", "🆚 Compare", "🌡️ Arrhenius", "⚡ Energy yield",
     "🔬 Imaging", "🔮 Forecast", "🧠 Report"])

# ------------------------------ DATA --------------------------------------
with tab_data:
    st.markdown("**Import outdoor / indoor aging data** - any CSV of "
                "timestamped IV parameters (+ environment). Map your "
                "columns once; the mapping is saved as a profile for next "
                "time.")
    up = st.file_uploader("Measurement CSV", type=["csv", "txt"],
                          key="pd_up")
    if up is not None:
        import io as _io
        raw_bytes = up.getvalue()
        # delimiter sniff
        head = raw_bytes[:4000].decode("utf-8", errors="replace")
        delim = ";" if head.count(";") > head.count(",") else ","
        if head.count("\t") > max(head.count(","), head.count(";")):
            delim = "\t"
        try:
            raw = _pd.read_csv(_io.BytesIO(raw_bytes), sep=delim)
        except Exception as e:
            st.error(f"Could not parse CSV: {e}")
            raw = None
        if raw is not None:
            st.caption(f"{len(raw):,} rows · {len(raw.columns)} columns "
                       f"(delimiter '{delim}')")
            st.dataframe(raw.head(8), use_container_width=True)
            profiles = load_profiles()
            prof_names = ["(new mapping)"] + sorted(profiles)
            prof_pick = st.selectbox("Import profile", prof_names,
                                     key="pd_prof")
            saved = (profiles.get(prof_pick, {}) if prof_pick in profiles
                     else {})
            saved_map = saved.get("colmap", {})
            cols = ["(none)"] + list(raw.columns)

            def _guess(canon):
                if saved_map.get(canon) in raw.columns:
                    return saved_map[canon]
                pats = {"timestamp": r"time|date", "device": r"dev|sample|module|id",
                        "pce": r"pce|eff", "pmax": r"pmax|power|p_max",
                        "isc": r"isc|i_sc", "voc": r"voc|v_oc",
                        "ff": r"^ff|fill", "imp": r"imp|i_mp",
                        "vmp": r"vmp|v_mp",
                        "gni": r"gni|poa|irr|ghi", "t_mod": r"t.?mod|tpv",
                        "t_amb": r"t.?amb|air", "rh": r"rh|humid",
                        "sweep": r"sweep|direction|scan",
                        "bias": r"bias|load"}
                pat = pats.get(canon)
                if pat:
                    for c in raw.columns:
                        if re.search(pat, str(c), re.I):
                            return c
                return "(none)"

            st.markdown("**Column mapping** (map at least the timestamp "
                        "and one performance metric):")
            colmap = {}
            grid = st.columns(3)
            for i, (canon, label) in enumerate(CANON_FIELDS):
                with grid[i % 3]:
                    g = _guess(canon)
                    sel = st.selectbox(label, cols,
                                       index=cols.index(g) if g in cols
                                       else 0,
                                       key=f"map_{canon}")
                    if sel != "(none)":
                        colmap[canon] = sel
            df1, df2, df3 = st.columns(3)
            with df1:
                dayfirst = st.checkbox("Day-first dates (31/12/2024)",
                                       value=bool(saved.get("dayfirst")),
                                       key="pd_dayfirst")
            with df2:
                ds_name = st.text_input("Save dataset as", key="pd_name",
                                        placeholder="e.g. ETL1_A_outdoor")
            with df3:
                prof_save = st.text_input("Save mapping as profile",
                                          key="pd_prof_name",
                                          value=(prof_pick
                                                 if prof_pick in profiles
                                                 else ""))
            if st.button("📥 Import dataset", type="primary", key="pd_imp",
                         disabled=not (ds_name.strip()
                                       and "timestamp" in colmap
                                       and any(m in colmap
                                               for m in METRIC_FIELDS))):
                try:
                    cdf = normalize_df(raw, colmap, dayfirst)
                except Exception as e:
                    st.error(f"Import failed: {e}")
                    cdf = None
                if cdf is not None and len(cdf):
                    PD_DATA_DIR.mkdir(parents=True, exist_ok=True)
                    safe = re.sub(r"[^A-Za-z0-9_-]", "_",
                                  ds_name.strip())[:60]
                    cdf.drop(columns=["date"]).to_csv(
                        PD_DATA_DIR / f"{safe}.csv", index=False)
                    if prof_save.strip():
                        profiles[prof_save.strip()] = {
                            "colmap": colmap, "dayfirst": dayfirst}
                        save_profiles(profiles)
                    st.success(f"Imported '{safe}': {len(cdf):,} rows, "
                               f"{cdf['device'].nunique()} device(s), "
                               f"{cdf['date'].nunique()} day(s), fields: "
                               + ", ".join(c for c in cdf.columns
                                           if c not in ("timestamp",
                                                        "date")))
                elif cdf is not None:
                    st.warning("No valid rows after parsing - check the "
                               "timestamp mapping / day-first setting.")

    if dsets:
        st.markdown("---")
        st.markdown("**Stored datasets**")
        for name in dsets:
            c1, c2 = st.columns([5, 1])
            c1.caption(name)
            if c2.button("🗑️", key=f"del_{name}"):
                (PD_DATA_DIR / f"{name}.csv").unlink(missing_ok=True)
                st.rerun()


# ------------------------------ I-V CURVES --------------------------------
with tab_iv:
    st.markdown("**Raw I-V sweeps** - upload curve files, get extracted "
                "parameters (Isc, Voc, FF, Pmax, PCE), hysteresis, and "
                "curve-evolution overlays. Extracted parameters can be "
                "saved as a dataset for the Long-term / Diurnal tabs.")
    iv_mode = st.radio("File layout",
                       ["Many files - one sweep per file (V, I columns)",
                        "One combined CSV - sweeps share the file"],
                       key="iv_mode")
    ic1, ic2 = st.columns(2)
    with ic1:
        iv_area = st.number_input("Active area (cm²)", 0.0, 10000.0,
                                  0.0, key="iv_area",
                                  help="Needed only for PCE; 0 = skip "
                                       "PCE.")
    with ic2:
        iv_sun = st.number_input("Irradiance during sweep (W/m²)",
                                 1.0, 1500.0, 1000.0, key="iv_sun")

    import io as _io

    def _read_curve_csv(data):
        head = data[:3000].decode("utf-8", errors="replace")
        delim = ";" if head.count(";") > head.count(",") else ","
        if head.count("\t") > max(head.count(","), head.count(";")):
            delim = "\t"
        raw = _pd.read_csv(_io.BytesIO(data), sep=delim,
                           engine="python", comment="#")
        return raw

    sweeps = []   # dicts: label, v, i, order
    if iv_mode.startswith("Many"):
        iv_files = st.file_uploader(
            "Curve files (each: one sweep)", type=["csv", "txt", "dat"],
            accept_multiple_files=True, key="iv_files")
        if iv_files:
            probe = _read_curve_csv(iv_files[0].getvalue())
            num_cols = [c for c in probe.columns
                        if _pd.to_numeric(probe[c], errors="coerce")
                        .notna().mean() > 0.7]
            vc1, vc2 = st.columns(2)
            with vc1:
                v_col = st.selectbox("Voltage column", num_cols,
                                     index=0, key="iv_vcol")
            with vc2:
                i_col = st.selectbox("Current column", num_cols,
                                     index=min(1, len(num_cols) - 1),
                                     key="iv_icol")
            for k, f in enumerate(sorted(iv_files,
                                         key=lambda x: x.name)):
                try:
                    raw = _read_curve_csv(f.getvalue())
                    v = _pd.to_numeric(raw[v_col], errors="coerce")
                    i = _pd.to_numeric(raw[i_col], errors="coerce")
                    sweeps.append({"label": f.name, "order": k,
                                   "v": v.values, "i": i.values})
                except Exception as e:
                    st.warning(f"{f.name}: {e}")
    else:
        iv_file = st.file_uploader("Combined CSV",
                                   type=["csv", "txt", "dat"],
                                   key="iv_file1")
        if iv_file is not None:
            raw = _read_curve_csv(iv_file.getvalue())
            cols = list(raw.columns)
            num_cols = [c for c in cols
                        if _pd.to_numeric(raw[c], errors="coerce")
                        .notna().mean() > 0.7]
            vc1, vc2, vc3 = st.columns(3)
            with vc1:
                v_col = st.selectbox("Voltage column", num_cols,
                                     key="iv_vcol2")
            with vc2:
                i_col = st.selectbox("Current column", num_cols,
                                     index=min(1, len(num_cols) - 1),
                                     key="iv_icol2")
            with vc3:
                g_col = st.selectbox("Sweep-ID / timestamp column",
                                     ["(none)"] + cols, key="iv_gcol")
            if g_col != "(none)":
                for k, (gid, g) in enumerate(raw.groupby(g_col,
                                                         sort=True)):
                    sweeps.append({
                        "label": str(gid), "order": k,
                        "v": _pd.to_numeric(g[v_col],
                                            errors="coerce").values,
                        "i": _pd.to_numeric(g[i_col],
                                            errors="coerce").values})
            else:
                sweeps.append({
                    "label": iv_file.name, "order": 0,
                    "v": _pd.to_numeric(raw[v_col],
                                        errors="coerce").values,
                    "i": _pd.to_numeric(raw[i_col],
                                        errors="coerce").values})

    if sweeps:
        rows_p = []
        for s in sweeps:
            prm = extract_iv_params(s["v"], s["i"])
            if prm is None:
                st.warning(f"'{s['label']}': unusable curve - skipped.")
                continue
            prm["sweep"] = sweep_direction(s["v"])
            prm["label"] = s["label"]
            prm["order"] = s["order"]
            if iv_area > 0:
                # assumes current in A; if your files are mA, PCE is
                # off by 1000x - obvious at a glance, then fix units.
                prm["pce"] = (prm["pmax"]
                              / (iv_sun * iv_area * 1e-4) * 100)
            rows_p.append(prm)
        if rows_p:
            pdf_ = _pd.DataFrame(rows_p)
            show_cols = [c for c in ("label", "sweep", "isc", "voc",
                                     "ff", "pmax", "vmp", "imp", "pce")
                         if c in pdf_.columns]
            st.dataframe(pdf_[show_cols].round(4),
                         use_container_width=True, height=280)
            # hysteresis: pair consecutive fwd/rev
            fwd = pdf_[pdf_["sweep"] == "fwd"]
            rev = pdf_[pdf_["sweep"] == "rev"]
            if len(fwd) and len(rev):
                m_f = fwd["pmax"].mean()
                m_r = rev["pmax"].mean()
                if m_r:
                    hi = (m_r - m_f) / m_r * 100
                    st.metric("Hysteresis index (mean Pmax, "
                              "(rev-fwd)/rev)", f"{hi:.1f}%")
            if alt is not None:
                longd = []
                for s in sweeps[:40]:
                    for vv, ii in zip(s["v"], s["i"]):
                        longd.append({"V": vv, "I": ii,
                                      "curve": s["label"][:30]})
                ld = _pd.DataFrame(longd).dropna()
                st.altair_chart(alt.Chart(ld).mark_line(
                    strokeWidth=1.5, opacity=0.8).encode(
                    x=alt.X("V:Q", title="Voltage (V)"),
                    y=alt.Y("I:Q", title="Current"),
                    color=alt.Color("curve:N", legend=None),
                    tooltip=["curve", "V", "I"]),
                    use_container_width=True)
                st.caption("First 40 curves shown; colour = sweep.")
            st.download_button(
                "⬇️ Extracted parameters CSV",
                data=pdf_.to_csv(index=False).encode("utf-8-sig"),
                file_name="iv_parameters.csv", mime="text/csv")


def _pick_dataset(key):
    ds = list_datasets()
    if not ds:
        st.info("Import a dataset in 📥 Data first.")
        return None, None, None
    c1, c2 = st.columns([2, 2])
    with c1:
        name = st.selectbox("Dataset", ds, key=f"{key}_ds")
    df = load_dataset(name)
    with c2:
        devs = sorted(df["device"].unique())
        dev = st.selectbox("Device", devs, key=f"{key}_dev")
    sub = df[df["device"] == dev]
    if "sweep" in sub.columns and sub["sweep"].nunique() > 1:
        sw = st.radio("Sweep", ["rev", "fwd", "both"], horizontal=True,
                      key=f"{key}_sw")
        if sw != "both":
            m = sub["sweep"].str.startswith(sw[0])
            if m.any():
                sub = sub[m]
    return name, dev, sub


# ------------------------------ LONG-TERM ---------------------------------
with tab_long:
    name, dev, sub = _pick_dataset("lt")
    if sub is not None and len(sub):
        avail = [m for m in METRIC_FIELDS if m in sub.columns
                 and sub[m].notna().any()]
        lc1, lc2 = st.columns(2)
        with lc1:
            metric = st.selectbox("Metric", avail, key="lt_metric")
        with lc2:
            pnom = st.number_input("Nameplate power Pnom (W, for PR)",
                                   0.0, 10000.0, 0.0, key="lt_pnom")
        daily = daily_series(sub, metric)
        if daily.empty:
            st.warning("No data for this metric.")
        else:
            fit = bilinear_fit(daily["day"], daily["norm"])
            k1, k2, k3, k4 = st.columns(4)
            k1.metric("Days of data", int(daily["day"].max()) + 1)
            t80_day = t80(daily)
            k2.metric("T80", f"day {t80_day}" if t80_day is not None
                      else "not reached")
            if fit:
                knee, a, b1, b2, fitted = fit
                k3.metric("Burn-in PLR", f"{b1 * 30 * 100:+.1f} %/mo")
                k4.metric(f"Stable PLR (knee d{knee:.0f})",
                          f"{b2 * 30 * 100:+.2f} %/mo")
                daily = daily.assign(fit=fitted)
                with st.expander("📐 Uncertainty on the fit "
                                 "(bootstrap 95% CI)"):
                    nb = st.slider("Resamples", 100, 500, 300, 50,
                                   key="lt_nb")
                    if st.button("Compute confidence intervals",
                                 key="lt_boot"):
                        with st.spinner(f"Refitting {nb} resampled "
                                        "datasets..."):
                            ci = bootstrap_bilinear(
                                daily["day"], daily["norm"],
                                n_boot=nb)
                        st.session_state["lt_ci"] = (
                            (name, dev, metric), ci)
                    stored = st.session_state.get("lt_ci")
                    ci = (stored[1] if stored and
                          stored[0] == (name, dev, metric) else None)
                    if ci:
                        u1, u2, u3, u4 = st.columns(4)
                        lo, hi = ci["burnin_plr_ci"]
                        u1.metric("Burn-in PLR 95% CI",
                                  f"[{lo:+.1f}, {hi:+.1f}] %/mo")
                        lo, hi = ci["stable_plr_ci"]
                        u2.metric("Stable PLR 95% CI",
                                  f"[{lo:+.2f}, {hi:+.2f}] %/mo")
                        lo, hi = ci["knee_ci"]
                        u3.metric("Knee 95% CI",
                                  f"[d{lo:.0f}, d{hi:.0f}]")
                        if ci["t80_ci"]:
                            lo, hi = ci["t80_ci"]
                            u4.metric("Model T80 95% CI",
                                      f"[d{lo:.0f}, d{hi:.0f}]")
                        else:
                            u4.metric("Model T80 95% CI",
                                      "not reached")
                        st.caption(
                            f"Pairs bootstrap, {ci['n_ok']}/"
                            f"{ci['n_boot']} successful refits. "
                            "Quote results as value [CI], e.g. stable "
                            "PLR "
                            f"{b2 * 3000:+.2f} [{ci['stable_plr_ci'][0]:+.2f}, "
                            f"{ci['stable_plr_ci'][1]:+.2f}] %/month. "
                            "T80 here is model-implied (where the "
                            "fitted curve crosses 0.80), so it can "
                            "differ from the empirical rolling-median "
                            "T80 above.")
                        st.session_state.setdefault(
                            "pd_metrics", {})["long_term_ci"] = {
                            "dataset": name, "device": dev,
                            "metric": metric, **{
                                kk: ci[kk] for kk in
                                ("burnin_plr_ci", "stable_plr_ci",
                                 "knee_ci", "t80_ci")}}
            plot = daily.rename(columns={"norm": "normalized"})
            if alt is not None:
                base = alt.Chart(plot).mark_circle(
                    size=30, opacity=0.6).encode(
                    x=alt.X("day:Q", title="Day of exposure"),
                    y=alt.Y("normalized:Q", title=f"Normalized {metric}",
                            scale=alt.Scale(zero=False)),
                    tooltip=["date:T", "value:Q", "normalized:Q"])
                layers = base
                if fit:
                    layers = base + alt.Chart(plot).mark_line(
                        color="#FF6B3D", strokeWidth=2.5).encode(
                        x="day:Q", y="fit:Q")
                st.altair_chart(layers, use_container_width=True)
            if pnom > 0:
                pr = perf_ratio_daily(sub, pnom)
                if not pr.empty:
                    st.markdown("**Performance ratio (daily)**")
                    if alt is not None:
                        st.altair_chart(alt.Chart(pr).mark_circle(
                            size=30, color="#8AB4F8",
                            opacity=0.6).encode(
                            x=alt.X("day:Q", title="Day"),
                            y=alt.Y("pr:Q", title="PR",
                                    scale=alt.Scale(zero=False)),
                            tooltip=["date:T", "pr:Q"]),
                            use_container_width=True)
            st.session_state.setdefault("pd_metrics", {})[
                "long_term"] = {
                "dataset": name, "device": dev, "metric": metric,
                "days": int(daily["day"].max()) + 1,
                "final_norm": round(float(
                    daily["norm"].tail(7).median()), 3),
                "t80_day": t80_day,
                "burnin_plr_pct_per_month": (round(b1 * 3000, 2)
                                             if fit else None),
                "stable_plr_pct_per_month": (round(b2 * 3000, 2)
                                             if fit else None),
                "knee_day": round(knee, 0) if fit else None,
            }
            dcol1, dcol2 = st.columns(2)
            dcol1.download_button(
                "⬇️ Daily series CSV",
                data=daily.to_csv(index=False).encode("utf-8-sig"),
                file_name=f"{name}_{dev}_{metric}_daily.csv",
                mime="text/csv")
            with dcol2:
                if st.button("🖼️ Export publication figure",
                             key="lt_fig"):
                    import matplotlib
                    matplotlib.use("Agg")
                    import matplotlib.pyplot as plt
                    fig_, ax = plt.subplots(figsize=(5.0, 3.3))
                    ax.plot(daily["day"], daily["norm"], "o", ms=2.5,
                            color="#666666", alpha=0.6,
                            label=f"{dev}")
                    if "fit" in daily.columns:
                        ax.plot(daily["day"], daily["fit"], "-",
                                color="#D2401E", lw=2,
                                label="bilinear fit")
                    _pub_ax(ax)
                    ax.set_xlabel("Day of exposure", fontsize=10)
                    ax.set_ylabel(f"Normalized {metric.upper()}",
                                  fontsize=10)
                    ax.legend(frameon=False, fontsize=8)
                    fig_.tight_layout()
                    png = save_pub_fig(fig_, f"longterm_{dev}_{metric}")
                    plt.close(fig_)
                    st.success(f"Saved: {png} (+ .svg)")
                    st.download_button("⬇️ PNG (300 dpi)",
                                       data=png.read_bytes(),
                                       file_name=png.name,
                                       mime="image/png",
                                       key="lt_fig_dl")

# ------------------------------ DIURNAL -----------------------------------
with tab_diurnal:
    name, dev, sub = _pick_dataset("di")
    if sub is not None and len(sub):
        if "gni" not in sub.columns or sub["gni"].isna().all():
            st.warning("The diurnal engine needs an irradiance column "
                       "(GNI/POA) - map it during import.")
        else:
            avail = [m for m in METRIC_FIELDS if m in sub.columns
                     and sub[m].notna().any()]
            d1, d2, d3, d4 = st.columns(4)
            with d1:
                dmetric = st.selectbox("Metric", avail, key="di_metric")
            with d2:
                g_c = st.number_input("Irradiance window centre (W/m²)",
                                      100, 1000, 400, step=50,
                                      key="di_g")
            with d3:
                g_w = st.number_input("± window (W/m²)", 10, 300, 50,
                                      step=10, key="di_w")
            with d4:
                dt_on = st.checkbox("ΔT ≤ 5 °C morning/evening",
                                    key="di_dt",
                                    help="Restricts to days where module "
                                         "temperature is similar at both "
                                         "points - per the paper's Vmp "
                                         "analysis.")
            dd = compute_dpd_dpr(sub, dmetric, g_c - g_w, g_c + g_w,
                                 dt_max=5 if dt_on else None)
            if dd.empty:
                st.warning("No valid days in this irradiance window - "
                           "widen it or check the data.")
            else:
                v1, v2, v3, v4 = st.columns(4)
                v1.metric("Valid days", len(dd))
                v2.metric("Median DPD", f"{dd['dpd'].median():.1f}%")
                v3.metric("Median DPR",
                          f"{dd['dpr'].median():.1f}%"
                          if dd["dpr"].notna().any() else "-")
                both = dd.dropna(subset=["dpr"])
                slope = None
                if len(both) >= 5:
                    import numpy as _np
                    slope = float(_np.polyfit(both["dpd"],
                                              both["dpr"], 1)[0])
                    v4.metric("DPR/DPD slope", f"{slope:.2f}")
                if alt is not None:
                    b1_, b2_ = st.columns(2)
                    with b1_:
                        st.markdown("**Monthly DPD**")
                        st.altair_chart(alt.Chart(dd).mark_boxplot(
                            color="#FF6B3D").encode(
                            x=alt.X("month:N", title="Month"),
                            y=alt.Y("dpd:Q", title="DPD (%)")),
                            use_container_width=True)
                    with b2_:
                        st.markdown("**Monthly DPR (overnight)**")
                        st.altair_chart(alt.Chart(
                            dd.dropna(subset=["dpr"])).mark_boxplot(
                            color="#8AB4F8").encode(
                            x=alt.X("month:N", title="Month"),
                            y=alt.Y("dpr:Q", title="DPR (%)")),
                            use_container_width=True)
                    c1_, c2_ = st.columns(2)
                    with c1_:
                        st.markdown("**DPR vs DPD** (linearity = "
                                    "reversible metastability)")
                        sc = alt.Chart(both).mark_circle(
                            size=50, opacity=0.6).encode(
                            x=alt.X("dpd:Q", title="DPD (%)"),
                            y=alt.Y("dpr:Q", title="DPR (%)"),
                            tooltip=["date:T", "dpd:Q", "dpr:Q"])
                        line = sc.transform_regression(
                            "dpd", "dpr").mark_line(color="#FF6B3D")
                        st.altair_chart(sc + line,
                                        use_container_width=True)
                    with c2_:
                        if "t_day" in dd.columns and \
                                dd["t_day"].notna().any():
                            st.markdown("**DPD by temperature bin**")
                            db = dd.dropna(subset=["t_day"]).copy()
                            db["Tbin"] = (db["t_day"] // 10 * 10
                                          ).astype(int).astype(str) \
                                + "-" + ((db["t_day"] // 10 * 10)
                                         .astype(int) + 10).astype(str)
                            st.altair_chart(alt.Chart(db).mark_boxplot(
                                color="#FFB08F").encode(
                                x=alt.X("Tbin:N", title="T (°C)"),
                                y=alt.Y("dpd:Q", title="DPD (%)")),
                                use_container_width=True)
                st.session_state.setdefault("pd_metrics", {})[
                    "diurnal"] = {
                    "dataset": name, "device": dev, "metric": dmetric,
                    "window_wm2": f"{g_c}±{g_w}",
                    "valid_days": len(dd),
                    "median_dpd_pct": round(float(dd["dpd"].median()), 2),
                    "median_dpr_pct": (round(float(dd["dpr"].median()), 2)
                                       if dd["dpr"].notna().any()
                                       else None),
                    "dpr_dpd_slope": (round(slope, 2)
                                      if slope is not None else None),
                }
                ddl1, ddl2 = st.columns(2)
                ddl1.download_button(
                    "⬇️ DPD/DPR series CSV",
                    data=dd.to_csv(index=False).encode("utf-8-sig"),
                    file_name=f"{name}_{dev}_dpd_dpr.csv",
                    mime="text/csv")
                with ddl2:
                    if st.button("🖼️ Export publication figure",
                                 key="di_fig"):
                        import matplotlib
                        matplotlib.use("Agg")
                        import matplotlib.pyplot as plt
                        months = sorted(dd["month"].unique())
                        fig_, axes = plt.subplots(
                            1, 2, figsize=(7.2, 3.2), sharey=True)
                        axes[0].boxplot(
                            [dd[dd["month"] == m]["dpd"].dropna()
                             for m in months], tick_labels=months)
                        axes[0].set_ylabel("DPD (%)", fontsize=10)
                        axes[1].boxplot(
                            [dd[dd["month"] == m]["dpr"].dropna()
                             for m in months], tick_labels=months)
                        axes[1].set_ylabel("DPR (%)", fontsize=10)
                        for ax_ in axes:
                            _pub_ax(ax_)
                            ax_.tick_params(axis="x", rotation=60,
                                            labelsize=7)
                        fig_.tight_layout()
                        png = save_pub_fig(fig_,
                                           f"diurnal_{dev}_{dmetric}")
                        plt.close(fig_)
                        st.success(f"Saved: {png} (+ .svg)")
                        st.download_button("⬇️ PNG (300 dpi)",
                                           data=png.read_bytes(),
                                           file_name=png.name,
                                           mime="image/png",
                                           key="di_fig_dl")

# ------------------------------ REVERSIBILITY -----------------------------
with tab_rev:
    st.markdown("**Reversible vs irreversible** - the decomposition that "
                "matters for perovskites: the morning-value trend is the "
                "irreversible envelope; the daily dip-and-recover cycle "
                "on top of it is the reversible metastability. Plus "
                "dark-storage recovery whenever the data has gaps.")
    name, dev, sub = _pick_dataset("rv")
    if sub is not None and len(sub):
        if "gni" not in sub.columns or sub["gni"].isna().all():
            st.warning("Needs an irradiance column (like the Diurnal "
                       "tab).")
        else:
            avail = [m for m in METRIC_FIELDS if m in sub.columns
                     and sub[m].notna().any()]
            rv1, rv2 = st.columns(2)
            with rv1:
                rmetric = st.selectbox("Metric", avail, key="rv_metric")
            with rv2:
                rg = st.number_input("Irradiance window (W/m², ±50)",
                                     100, 1000, 400, step=50,
                                     key="rv_g")
            dd = compute_dpd_dpr(sub, rmetric, rg - 50, rg + 50)
            if dd.empty or len(dd) < 5:
                st.warning("Not enough valid days.")
            else:
                summ, gaps = day_night_analysis(dd)
                if summ:
                    s1, s2, s3, s4 = st.columns(4)
                    s1.metric("Irreversible rate",
                              f"{summ['irreversible_rate_pct_per_month']:+.2f} %/mo")
                    s2.metric("Reversible amplitude",
                              f"{summ['reversible_amplitude_pct']:.1f}%")
                    s3.metric("Recovery efficiency",
                              f"{summ['recovery_efficiency_pct']:.0f}%"
                              if summ["recovery_efficiency_pct"]
                              is not None else "-")
                    s4.metric("Total morning loss",
                              f"{summ['total_morning_loss_pct']:.1f}%")
                    if alt is not None:
                        pl = dd.melt(id_vars=["date"],
                                     value_vars=["morning", "evening"],
                                     var_name="point",
                                     value_name="value")
                        st.altair_chart(alt.Chart(pl).mark_circle(
                            size=35, opacity=0.65).encode(
                            x=alt.X("date:T", title="Date"),
                            y=alt.Y("value:Q", title=rmetric,
                                    scale=alt.Scale(zero=False)),
                            color=alt.Color(
                                "point:N",
                                scale=alt.Scale(
                                    domain=["morning", "evening"],
                                    range=["#8AB4F8", "#FF6B3D"])),
                            tooltip=["date:T", "point:N", "value:Q"]),
                            use_container_width=True)
                        st.caption("Blue mornings drifting down = "
                                   "irreversible; the blue-orange "
                                   "daily split = reversible.")
                    if not gaps.empty:
                        st.markdown("**🌑 Dark-storage / outage "
                                    "recovery** - performance across "
                                    "data gaps > 2 days:")
                        st.dataframe(gaps, use_container_width=True,
                                     hide_index=True)
                        st.caption("Positive recovery across a gap = "
                                   "the reversible component healing "
                                   "during storage - your paper's "
                                   "indoor-recovery experiment, found "
                                   "automatically.")
                    st.session_state.setdefault("pd_metrics", {})[
                        "reversibility"] = {
                        "dataset": name, "device": dev,
                        "metric": rmetric, **summ,
                        "storage_gaps_found": len(gaps)}

# ------------------------------ COMPARE -----------------------------------
def _all_pairs():
    pairs = []
    for nm in list_datasets():
        try:
            d = load_dataset(nm)
            for dv in sorted(d["device"].unique()):
                pairs.append(f"{nm} · {dv}")
        except Exception:
            continue
    return pairs


def _pair_daily(pair, metric):
    nm, dv = pair.split(" · ", 1)
    d = load_dataset(nm)
    sub = d[d["device"] == dv]
    if "sweep" in sub.columns and sub["sweep"].astype(str) \
            .str.startswith("r").any():
        sub = sub[sub["sweep"].astype(str).str.startswith("r")]
    if metric not in sub.columns or sub[metric].isna().all():
        return None
    return daily_series(sub, metric)


with tab_cmp:
    st.markdown("**Device comparison** - encapsulant A vs B, ETL1 vs "
                "ETL2: overlaid normalized decay plus the figures of "
                "merit side by side.")
    pairs = _all_pairs()
    if not pairs:
        st.info("Import datasets first.")
    else:
        cm1, cm2 = st.columns([3, 1])
        with cm1:
            cmp_pick = st.multiselect("Devices (2-6)", pairs,
                                      max_selections=6, key="cmp_pick")
        with cm2:
            cmp_metric = st.selectbox("Metric", METRIC_FIELDS,
                                      key="cmp_metric")
        if len(cmp_pick) >= 2:
            series, rows = [], []
            for p in cmp_pick:
                dl = _pair_daily(p, cmp_metric)
                if dl is None or dl.empty:
                    st.warning(f"{p}: no {cmp_metric} data.")
                    continue
                fit = bilinear_fit(dl["day"], dl["norm"])
                row = {"Device": p,
                       "Days": int(dl["day"].max()) + 1,
                       "T80": t80(dl),
                       "Final (norm.)": round(float(
                           dl["norm"].tail(7).median()), 3)}
                if fit:
                    knee, a, b1, b2, fitted = fit
                    row["Burn-in %/mo"] = round(b1 * 3000, 2)
                    row["Stable %/mo"] = round(b2 * 3000, 2)
                    row["Knee day"] = round(knee)
                rows.append(row)
                dl = dl.assign(devlabel=p)
                series.append(dl)
            if series:
                alld = _pd.concat(series)
                if alt is not None:
                    st.altair_chart(alt.Chart(alld).mark_circle(
                        size=22, opacity=0.55).encode(
                        x=alt.X("day:Q", title="Day of exposure"),
                        y=alt.Y("norm:Q",
                                title=f"Normalized {cmp_metric}",
                                scale=alt.Scale(zero=False)),
                        color=alt.Color("devlabel:N", title="Device"),
                        tooltip=["devlabel", "date:T", "norm:Q"]),
                        use_container_width=True)
                st.dataframe(_pd.DataFrame(rows),
                             use_container_width=True,
                             hide_index=True)
                st.caption("Stable-phase rates within ~±0.3 %/mo of "
                           "each other are usually not distinguishable "
                           "on noisy outdoor data - treat close calls "
                           "as ties.")
                if st.button("🖼️ Export publication figure",
                             key="cmp_fig"):
                    import matplotlib
                    matplotlib.use("Agg")
                    import matplotlib.pyplot as plt
                    fig, ax = plt.subplots(figsize=(5.2, 3.4))
                    for p in sorted(alld["devlabel"].unique()):
                        s = alld[alld["devlabel"] == p]
                        ax.plot(s["day"], s["norm"], "o", ms=2.5,
                                alpha=0.6, label=p)
                    _pub_ax(ax)
                    ax.set_xlabel("Day of exposure", fontsize=10)
                    ax.set_ylabel(f"Normalized {cmp_metric.upper()}",
                                  fontsize=10)
                    ax.legend(frameon=False, fontsize=8)
                    fig.tight_layout()
                    png = save_pub_fig(fig, "compare_" + cmp_metric)
                    plt.close(fig)
                    st.success(f"Saved: {png} (+ .svg)")
                    st.download_button("⬇️ PNG (300 dpi)",
                                       data=png.read_bytes(),
                                       file_name=png.name,
                                       mime="image/png",
                                       key="cmp_fig_dl")

# ------------------------------ ARRHENIUS ---------------------------------
with tab_arr:
    st.markdown("**Arrhenius analysis** - degradation rates from "
                "devices aged at different temperatures → activation "
                "energy and acceleration factors. Needs ≥3 devices at "
                "distinct temperatures.")
    pairs = _all_pairs()
    if not pairs:
        st.info("Import datasets first.")
    else:
        ar_pick = st.multiselect("Devices", pairs, key="ar_pick")
        ar_metric = st.selectbox("Metric", METRIC_FIELDS,
                                 key="ar_metric")
        ar_src = st.radio("Rate to use",
                          ["Stable-phase PLR (post burn-in)",
                           "Overall linear slope"],
                          horizontal=True, key="ar_src")
        if len(ar_pick) >= 3:
            tbl = _pd.DataFrame({
                "Device": ar_pick,
                "Aging temperature (°C)": [None] * len(ar_pick)})
            ttbl = st.data_editor(tbl, use_container_width=True,
                                  hide_index=True, key="ar_tbl")
            if st.button("🌡️ Fit Arrhenius", type="primary",
                         key="ar_go"):
                temps, rates, used = [], [], []
                import numpy as _np
                for _, r in ttbl.iterrows():
                    tc = _pd.to_numeric(
                        _pd.Series([r["Aging temperature (°C)"]]),
                        errors="coerce").iloc[0]
                    if _pd.isna(tc):
                        continue
                    dl = _pair_daily(r["Device"], ar_metric)
                    if dl is None or dl.empty:
                        continue
                    fit = bilinear_fit(dl["day"], dl["norm"])
                    if ar_src.startswith("Stable") and fit:
                        rate = -fit[3] * 3000  # %/mo, positive = loss
                    else:
                        rate = -float(_np.polyfit(
                            dl["day"], dl["norm"], 1)[0]) * 3000
                    if rate <= 0:
                        st.warning(f"{r['Device']}: non-degrading rate "
                                   f"({rate:.2f} %/mo) - excluded.")
                        continue
                    temps.append(float(tc))
                    rates.append(rate)
                    used.append({"Device": r["Device"], "T (°C)": tc,
                                 "Rate (%/mo)": round(rate, 3)})
                res = arrhenius_fit(temps, rates)
                if res is None:
                    st.error("Need ≥3 devices with positive rates and "
                             "temperatures.")
                else:
                    ea, lna, r2 = res
                    a1, a2, a3 = st.columns(3)
                    a1.metric("Activation energy", f"{ea:.2f} eV")
                    a2.metric("R²", f"{r2:.3f}")
                    af = acceleration_factor(ea, 25, 65)
                    a3.metric("AF (25→65 °C)", f"{af:.1f}×")
                    st.dataframe(_pd.DataFrame(used),
                                 use_container_width=True,
                                 hide_index=True)
                    import numpy as _np
                    pl = _pd.DataFrame({
                        "invT": [1000 / (t + 273.15) for t in temps],
                        "lnk": [_np.log(r) for r in rates]})
                    if alt is not None:
                        sc = alt.Chart(pl).mark_circle(
                            size=90, color="#FF6B3D").encode(
                            x=alt.X("invT:Q", title="1000/T (1/K)",
                                    scale=alt.Scale(zero=False)),
                            y=alt.Y("lnk:Q", title="ln(rate)",
                                    scale=alt.Scale(zero=False)))
                        st.altair_chart(
                            sc + sc.transform_regression(
                                "invT", "lnk").mark_line(
                                color="#8AB4F8"),
                            use_container_width=True)
                    st.caption("⚠ Arrhenius extrapolation assumes ONE "
                               "thermally activated mechanism across "
                               "the range - perovskites can switch "
                               "mechanisms with temperature. Treat AF "
                               "values as first-order estimates.")
        else:
            st.caption("Pick at least 3 devices, then enter each one's "
                       "aging temperature.")

# ------------------------------ ENERGY YIELD ------------------------------
with tab_yield:
    st.markdown("**Energy yield** - the number a customer or investor "
                "actually cares about.")
    y_mode = st.radio("Mode", ["Measured (from a dataset)",
                               "Projected (location model)"],
                      horizontal=True, key="y_mode")

    if y_mode.startswith("Measured"):
        name, dev, sub = _pick_dataset("ey")
        if sub is not None and len(sub):
            if "pmax" not in sub.columns or sub["pmax"].isna().all():
                st.warning("This dataset has no power column (pmax).")
            else:
                ey_pnom = st.number_input(
                    "Nameplate power Pnom (W)", 0.0, 100000.0, 0.0,
                    key="ey_pnom",
                    help="Enables specific yield (kWh/kWp) and PR.")
                de = daily_energy(sub, ey_pnom or None)
                if de.empty:
                    st.warning("Not enough samples per day to "
                               "integrate.")
                else:
                    tot_kwh = de["energy_wh"].sum() / 1000
                    y1, y2, y3, y4 = st.columns(4)
                    y1.metric("Days integrated", len(de))
                    y2.metric("Total energy",
                              f"{tot_kwh * 1000:.1f} Wh"
                              if tot_kwh < 1 else f"{tot_kwh:.2f} kWh")
                    if ey_pnom > 0:
                        y3.metric("Specific yield",
                                  f"{tot_kwh / (ey_pnom / 1000):.0f} "
                                  "kWh/kWp")
                    if "pr" in de.columns and de["pr"].notna().any():
                        y4.metric("Mean PR",
                                  f"{de['pr'].mean():.2f}")
                    mo = (de.groupby("month")["energy_wh"].sum()
                          .reset_index())
                    mo["kWh"] = mo["energy_wh"] / 1000
                    if alt is not None:
                        st.altair_chart(alt.Chart(mo).mark_bar(
                            color="#FF6B3D").encode(
                            x=alt.X("month:N", title="Month"),
                            y=alt.Y("kWh:Q", title="Energy (kWh)"),
                            tooltip=["month", "kWh"]),
                            use_container_width=True)
                    st.caption("⚠ Integration covers only the hours "
                               "actually sampled (mean "
                               f"{de['hours'].mean():.1f} h/day here) - "
                               "loggers that skip low light "
                               "underestimate true yield.")
                    st.download_button(
                        "⬇️ Daily energy CSV",
                        data=de.to_csv(index=False)
                        .encode("utf-8-sig"),
                        file_name=f"{name}_{dev}_energy.csv",
                        mime="text/csv")

    else:
        st.caption("First-order model: E = Pstc × H_POA/1kW × PR × "
                   "temperature correction. Good for comparisons and "
                   "pitch material; not a substitute for full "
                   "simulation.")
        p1, p2, p3 = st.columns(3)
        with p1:
            yp_p = st.number_input("Module rated power Pstc (W)",
                                   0.01, 1e6, 400.0, key="yp_p")
        with p2:
            yp_pr = st.number_input("Performance ratio", 0.4, 1.0,
                                    0.80, step=0.01, key="yp_pr")
        with p3:
            yp_tc = st.number_input("Temp. coefficient (%/°C)",
                                    -1.0, 0.0, -0.30, step=0.01,
                                    key="yp_tc",
                                    help="Perovskites are typically "
                                         "gentler than c-Si (-0.4): "
                                         "-0.2 to -0.3 %/°C reported.")
        p4, p5 = st.columns(2)
        with p4:
            preset = st.selectbox("Insolation preset",
                                  list(YIELD_PRESETS.keys()),
                                  key="yp_preset")
        with p5:
            yp_deg = st.number_input("Degradation (%/year)", 0.0, 30.0,
                                     1.0, step=0.5, key="yp_deg")
        base_tbl = _pd.DataFrame({
            "Month": MONTHS,
            "POA insolation (kWh/m²)": YIELD_PRESETS[preset],
            "Avg ambient T (°C)": [12, 13, 15, 19, 24, 28, 31, 31, 28,
                                   23, 17, 13],
        })
        tbl = st.data_editor(base_tbl, use_container_width=True,
                             hide_index=True, key=f"yp_tbl_{preset}")
        h = _pd.to_numeric(tbl["POA insolation (kWh/m²)"],
                           errors="coerce").fillna(0)
        ta = _pd.to_numeric(tbl["Avg ambient T (°C)"],
                            errors="coerce").fillna(20)
        tcell = ta + 20  # simple NOCT-style offset
        tfac = 1 + yp_tc / 100 * (tcell - 25)
        e_month = yp_p / 1000 * h * yp_pr * tfac  # kWh
        annual = float(e_month.sum())
        m1, m2 = st.columns(2)
        m1.metric("Year-1 energy", f"{annual:,.0f} kWh")
        m2.metric("Specific yield",
                  f"{annual / (yp_p / 1000):,.0f} kWh/kWp")
        res = _pd.DataFrame({"Month": MONTHS,
                             "kWh": e_month.round(1)})
        if alt is not None:
            st.altair_chart(alt.Chart(res).mark_bar(
                color="#8AB4F8").encode(
                x=alt.X("Month:N", sort=MONTHS),
                y=alt.Y("kWh:Q"),
                tooltip=["Month", "kWh"]),
                use_container_width=True)
        yrs = _pd.DataFrame({
            "Year": list(range(1, 11)),
            "kWh": [round(annual * (1 - yp_deg / 100) ** (n - 0.5), 1)
                    for n in range(1, 11)]})
        yrs["Cumulative kWh"] = yrs["kWh"].cumsum().round(0)
        st.markdown("**10-year projection** (with the degradation rate "
                    "above)")
        st.dataframe(yrs, use_container_width=True, hide_index=True,
                     height=240)
        st.caption("⚠ Model assumptions: NOCT-style +20 °C cell "
                   "offset, constant PR, linear-compounded "
                   "degradation. Label any use of these numbers as "
                   "estimates.")

# ------------------------------ IMAGING -----------------------------------
with tab_img:
    st.markdown("**Imaging suite** - PL, EL and DLIT series over aging: "
                "intensity decay, dark-area growth, hotspot detection, "
                "and first-vs-last difference maps. Upload images in "
                "time order (or with sortable names like "
                "`M1_PL_2025-03-01.tif`).")
    st.caption("⚠ Metrics compare images to each other - keep camera "
               "settings (exposure, gain, current/excitation) constant "
               "across the series, or the trends measure your settings, "
               "not your device.")
    im1, im2 = st.columns(2)
    with im1:
        modality = st.selectbox("Modality", ["PL", "EL", "DLIT"],
                                key="img_mod")
    with im2:
        dark_ratio = st.slider(
            "Dark-area threshold (× reference median)", 0.1, 0.9, 0.5,
            step=0.05, key="img_thr",
            help="A pixel counts as 'dark/degraded' below this fraction "
                 "of the FIRST image's median intensity.")
    img_files = st.file_uploader(
        "Image series (PNG / TIFF / JPG - 8 or 16 bit)",
        type=["png", "tif", "tiff", "jpg", "jpeg", "bmp"],
        accept_multiple_files=True, key="img_files")
    if img_files:
        try:
            from PIL import Image  # noqa: F401
        except ImportError:
            st.error("Pillow is missing: pip install pillow")
            img_files = []
    if img_files:
        files = sorted(img_files, key=lambda f: f.name)
        arrs, stats = [], []
        ref_median = None
        bar = st.progress(0.0)
        for k, f in enumerate(files, start=1):
            bar.progress(k / len(files), text=f.name)
            try:
                a = load_image_gray(f.getvalue())
            except Exception as e:
                st.warning(f"{f.name}: {e}")
                continue
            if ref_median is None:
                ref_median = float(__import__("numpy").median(a))
            s = image_stats(a, ref_median, dark_ratio)
            s["image"] = f.name
            s["idx"] = k - 1
            stats.append(s)
            arrs.append(a)
        bar.empty()
        if stats:
            sdf = _pd.DataFrame(stats)
            sdf["mean_norm"] = sdf["mean"] / sdf["mean"].iloc[0]
            key_metric = ("hot_frac" if modality == "DLIT"
                          else "dark_frac")
            g1, g2, g3 = st.columns(3)
            g1.metric("Images", len(sdf))
            g2.metric("Mean intensity, last vs first",
                      f"{sdf['mean_norm'].iloc[-1] * 100:.0f}%")
            lbl = ("Hotspot area (last)" if modality == "DLIT"
                   else "Dark area (last)")
            g3.metric(lbl,
                      f"{sdf[key_metric].iloc[-1] * 100:.1f}%")
            if alt is not None:
                mvis = sdf.melt(
                    id_vars=["idx", "image"],
                    value_vars=["mean_norm", key_metric],
                    var_name="metric", value_name="value")
                st.altair_chart(alt.Chart(mvis).mark_line(
                    point=True).encode(
                    x=alt.X("idx:Q", title="Image # (time order)"),
                    y=alt.Y("value:Q", title="Value"),
                    color=alt.Color("metric:N"),
                    tooltip=["image", "metric", "value"]),
                    use_container_width=True)
            st.dataframe(sdf[["image", "mean", "median",
                              "dark_frac" if "dark_frac" in sdf
                              else "hot_frac", "hot_frac",
                              "mean_norm"]].round(4),
                         use_container_width=True, height=240)
            if len(arrs) >= 2:
                import numpy as _np
                a0, a1 = arrs[0], arrs[-1]
                h = min(a0.shape[0], a1.shape[0])
                w = min(a0.shape[1], a1.shape[1])
                a0, a1 = a0[:h, :w], a1[:h, :w]

                def _to8(a):
                    lo, hi = _np.percentile(a, [1, 99])
                    return (255 * _np.clip((a - lo) / max(hi - lo, 1e-9),
                                           0, 1)).astype("uint8")
                d1, d2, d3 = st.columns(3)
                d1.image(_to8(a0), caption="First",
                         use_container_width=True)
                d2.image(_to8(a1), caption="Last",
                         use_container_width=True)
                diff = a1 - a0
                d3.image(_to8(diff), caption="Difference (last-first)",
                         use_container_width=True)
                st.caption("In the difference map: dark = intensity "
                           "lost (PL/EL degradation), bright = gained "
                           "(DLIT: new heat).")
            st.download_button(
                "⬇️ Image metrics CSV",
                data=sdf.to_csv(index=False).encode("utf-8-sig"),
                file_name=f"{modality}_series_metrics.csv",
                mime="text/csv")
            st.session_state.setdefault("pd_metrics", {})[
                "imaging_" + modality] = {
                "images": len(sdf),
                "mean_intensity_last_vs_first_pct":
                    round(float(sdf["mean_norm"].iloc[-1] * 100), 1),
                ("hotspot_area_last_pct" if modality == "DLIT"
                 else "dark_area_last_pct"):
                    round(float(sdf[key_metric].iloc[-1] * 100), 2)}

# ------------------------------ FORECAST ----------------------------------
with tab_forecast:
    name, dev, sub = _pick_dataset("fc")
    if sub is not None and len(sub):
        need = [c for c in ("pmax", "gni") if c not in sub.columns
                or sub[c].isna().all()]
        if need:
            st.warning(f"Forecast needs pmax and gni columns "
                       f"(missing: {', '.join(need)}).")
        else:
            feats = ["gni"] + [c for c in ("t_mod", "t_amb")
                               if c in sub.columns
                               and sub[c].notna().any()][:1]
            st.caption(f"Model: power = f({', '.join(feats)}), "
                       "70/30 random split - after the XGBoost workflow "
                       "in your ACS Energy Lett. paper.")
            if st.button("🔮 Train & evaluate", type="primary",
                         key="fc_go"):
                w = sub.dropna(subset=["pmax"] + feats)
                w = w[w["gni"] > 50]
                if len(w) < 200:
                    st.warning("Fewer than 200 daylight samples - not "
                               "enough to train.")
                else:
                    X = w[feats].values
                    y = w["pmax"].values
                    try:
                        from xgboost import XGBRegressor
                        model_ml = XGBRegressor(n_estimators=300,
                                                max_depth=5,
                                                learning_rate=0.08,
                                                verbosity=0)
                        used = "XGBoost"
                    except ImportError:
                        try:
                            from sklearn.ensemble import (
                                GradientBoostingRegressor)
                            model_ml = GradientBoostingRegressor(
                                n_estimators=300, max_depth=4)
                            used = "sklearn GradientBoosting"
                        except ImportError:
                            model_ml = None
                            st.error("Install once: pip install xgboost "
                                     "scikit-learn")
                    if model_ml is not None:
                        import numpy as _np
                        rng = _np.random.default_rng(42)
                        idx = rng.permutation(len(w))
                        cut = int(len(w) * 0.7)
                        tr, te = idx[:cut], idx[cut:]
                        model_ml.fit(X[tr], y[tr])
                        pred = model_ml.predict(X[te])
                        nr, nm = nrmse_nmbe(y[te], pred)
                        f1, f2, f3 = st.columns(3)
                        f1.metric("Model", used)
                        f2.metric("nRMSE", f"{nr:.2f}%")
                        f3.metric("nMBE", f"{nm:+.2f}%")
                        te_df = w.iloc[te].copy()
                        te_df["predicted"] = pred
                        if alt is not None:
                            st.altair_chart(alt.Chart(
                                te_df.sample(min(3000, len(te_df)),
                                             random_state=1)
                            ).mark_circle(size=25, opacity=0.4).encode(
                                x=alt.X("pmax:Q", title="Measured power"),
                                y=alt.Y("predicted:Q",
                                        title="Predicted power"),
                                tooltip=["timestamp:T", "pmax:Q",
                                         "predicted:Q"]),
                                use_container_width=True)
                        st.session_state["fc_result"] = te_df
                        st.session_state.setdefault("pd_metrics", {})[
                            "forecast"] = {
                            "dataset": name, "device": dev,
                            "model": used, "features": feats,
                            "nrmse_pct": round(nr, 2),
                            "nmbe_pct": round(nm, 2)}
            te_df = st.session_state.get("fc_result")
            if te_df is not None:
                days_av = sorted({str(d) for d in te_df["date"]})
                dpick = st.selectbox("Day overlay (test set)", days_av,
                                     key="fc_day")
                dday = te_df[te_df["date"].astype(str) == dpick] \
                    .sort_values("timestamp")
                if alt is not None and not dday.empty:
                    m1 = alt.Chart(dday).mark_line(
                        color="#8AB4F8").encode(
                        x="timestamp:T",
                        y=alt.Y("pmax:Q", title="Power"))
                    m2 = alt.Chart(dday).mark_line(
                        color="#FF6B3D",
                        strokeDash=[5, 3]).encode(
                        x="timestamp:T", y="predicted:Q")
                    st.altair_chart(m1 + m2, use_container_width=True)
                    st.caption("Blue: measured · Orange dashed: "
                               "predicted")

# ------------------------------ REPORT ------------------------------------
with tab_report:
    st.markdown("**Degradation report** - turns the metrics computed in "
                "the other tabs into a structured analysis: performance "
                "summary, diurnal behaviour, mechanism hypotheses, and "
                "recommended next measurements.")
    metrics = st.session_state.get("pd_metrics", {})
    if not metrics:
        st.info("Run analyses first (Long-term / Diurnal / Forecast) - "
                "their results feed this report.")
    else:
        import json as _json
        st.json(metrics)
        r1, r2, r3 = st.columns([1.5, 1.5, 2])
        with r1:
            rep_model_label = st.selectbox("Model", list(MODELS.keys()),
                                           index=3, key="rp_model")
        with r2:
            ground = st.checkbox("Ground in paper library", value=True,
                                 key="rp_ground",
                                 help="Retrieves mechanism literature "
                                      "from your corpus (first use loads "
                                      "the embedding model, ~15 s).")
        with r3:
            rp_notes = st.text_input("Context notes (optional)",
                                     key="rp_notes",
                                     placeholder="e.g. p-i-n mini-module, "
                                                 "NiO/perovskite/LiF-C60, "
                                                 "Nicosia outdoor")
        if st.button("🧠 Generate report", type="primary", key="rp_go"):
            can_call = (st.session_state.get("backend") == "max"
                        or bool(api_key.strip()))
            if not can_call:
                st.error("Needs the API key (sidebar).")
            else:
                excerpts = ""
                hits_n = 0
                if ground:
                    try:
                        with st.spinner("Retrieving mechanism "
                                        "literature..."):
                            import chromadb
                            import config
                            from chromadb.utils.embedding_functions \
                                import SentenceTransformerEmbeddingFunction
                            ef = SentenceTransformerEmbeddingFunction(
                                model_name=config.EMBED_MODEL)
                            col = chromadb.PersistentClient(
                                path=str(config.DB_DIR)).get_collection(
                                config.COLLECTION_NAME,
                                embedding_function=ef)
                            q = ("perovskite outdoor degradation diurnal "
                                 "recovery ion migration burn-in "
                                 "reversible metastability "
                                 + (rp_notes or ""))
                            res = col.query(query_texts=[q], n_results=10,
                                            include=["documents",
                                                     "metadatas"])
                            lines = []
                            for d, m in zip(res["documents"][0],
                                            res["metadatas"][0]):
                                lines.append(
                                    f"[{len(lines) + 1}] "
                                    f"{m.get('title', '?')}: {d[:400]}")
                            hits_n = len(lines)
                            excerpts = ("\n\nLITERATURE EXCERPTS (cite "
                                        "as [n]):\n" + "\n\n".join(lines))
                    except Exception as e:
                        st.warning(f"Corpus grounding unavailable ({e}) "
                                   "- writing ungrounded report.")
                umsg = (f"DEVICE CONTEXT: {rp_notes or 'not specified'}"
                        f"\n\nCOMPUTED METRICS:\n"
                        f"{_json.dumps(metrics, indent=1)}{excerpts}")
                with st.spinner(f"{rep_model_label} is writing..."):
                    try:
                        out = call_claude(api_key.strip(),
                                          DEG_REPORT_SYSTEM, umsg,
                                          MODELS[rep_model_label],
                                          max_tokens=6000)
                    except Exception as e:
                        st.error(f"Claude error: {e}")
                        out = None
                if out:
                    st.session_state["last_degreport"] = out
                    save_report(metrics.get("long_term", {})
                                .get("device", "device"), out)
                    if hits_n:
                        st.caption(f"Grounded in {hits_n} corpus "
                                   "passages.")
        if st.session_state.get("last_degreport"):
            st.markdown("---")
            st.markdown(st.session_state["last_degreport"])
            st.download_button(
                "⬇️ Report as Markdown",
                data=st.session_state["last_degreport"].encode("utf-8"),
                file_name="PeroDeg_report.md", mime="text/markdown")
            st.caption("Also auto-saved to answers/ as formatted .docx "
                       "+ .md.")

        st.markdown("---")
        st.markdown("**📋 ISOS summary table** - the consensus-style "
                    "reporting block (protocol, conditions, figures of "
                    "merit) built deterministically from the computed "
                    "metrics - ready for a manuscript's SI.")
        i1, i2 = st.columns(2)
        with i1:
            isos_proto = st.selectbox(
                "Protocol", ["ISOS-O-1", "ISOS-O-2", "ISOS-O-3",
                             "ISOS-L-1", "ISOS-L-2", "ISOS-L-3",
                             "ISOS-D-1", "ISOS-D-2", "ISOS-D-3",
                             "ISOS-LC-1", "ISOS-T-1", "custom"],
                index=2, key="isos_proto",
                help="O = outdoor, L = light soaking, D = dark storage, "
                     "LC = light cycling, T = thermal.")
            isos_bias = st.text_input("Bias between measurements",
                                      value="open circuit",
                                      key="isos_bias")
        with i2:
            isos_sample = st.text_input(
                "Sample description", key="isos_sample",
                placeholder="e.g. p-i-n mini-module, 4 cm², "
                            "NiO/FACsPb(I,Br)3/LiF-C60, laminated")
            isos_init = st.text_input("Initial performance",
                                      key="isos_init",
                                      placeholder="e.g. PCE 14.3% "
                                                  "(reverse, STC)")
        isos_loc = st.text_input("Location / setup", key="isos_loc",
                                 placeholder="e.g. fixed-tilt outdoor "
                                             "array, Nicosia, CY; IV "
                                             "every 15 min at GNI>400")
        if st.button("📋 Build ISOS table", key="isos_go"):
            table_md = isos_table(
                {"protocol": isos_proto, "sample": isos_sample,
                 "location": isos_loc, "bias": isos_bias,
                 "initial": isos_init}, metrics)
            st.session_state["isos_md"] = table_md
            save_report("ISOS_summary", table_md)
        if st.session_state.get("isos_md"):
            st.markdown(st.session_state["isos_md"])
            st.caption("Saved to answers/ as formatted .docx + .md "
                       "(real Word table).")
