"""
PV Radar - by GrapheAI. Standalone PV / perovskite news intelligence.

Run with:
    streamlit run pvradar.py --server.port 8502

Lives in the same folder as GrapheAI (workbench.py) and shares:
  - answers/pv_radar.json   the accumulated news store (full history)
  - answers/spend.json      the monthly API spend + budget
  - the Claude backend      API key or Claude Max via Claude Code

It does NOT load the paper index or embedding model, so it starts in
seconds and can run side by side with GrapheAI (different port).
"""

import datetime
import re
from pathlib import Path

import streamlit as st

ANSWERS_DIR = Path("answers")

MODELS = {
    "Fast (Haiku 4.5)": "claude-haiku-4-5",
    "Balanced (Sonnet 5)": "claude-sonnet-5",
    "Deep (Opus 5)": "claude-opus-5",
    "Frontier (Fable 5.1)": "claude-fable-5-1",
}

# API list prices, $ per million tokens (input, output) - for the sidebar
# cost estimate. Update if Anthropic's pricing changes.
PRICES = {
    "claude-haiku-4-5": (1.00, 5.00),
    "claude-sonnet-4-6": (3.00, 15.00),
    "claude-sonnet-5": (2.00, 10.00),
    "claude-opus-4-8": (5.00, 25.00),
    "claude-opus-5": (5.00, 25.00),
    "claude-fable-5": (10.00, 50.00),
    "claude-fable-5-1": (10.00, 50.00),
}

SPEND_FILE = ANSWERS_DIR / "spend.json"


def _load_spend():
    import json
    if SPEND_FILE.exists():
        try:
            return json.loads(SPEND_FILE.read_text(encoding="utf-8"))
        except Exception:
            return {}
    return {}


def _write_spend(data):
    import json
    try:
        ANSWERS_DIR.mkdir(exist_ok=True)
        SPEND_FILE.write_text(json.dumps(data), encoding="utf-8")
    except Exception:
        pass


def _add_spend(amount):
    if amount <= 0:
        return
    data = _load_spend()
    key = datetime.date.today().strftime("%Y-%m")
    data[key] = round(float(data.get(key, 0.0)) + amount, 5)
    _write_spend(data)


def month_spend():
    return float(_load_spend().get(datetime.date.today().strftime("%Y-%m"),
                                   0.0))


def get_budget():
    return float(_load_spend().get("_budget", 0.0))


def set_budget(value):
    data = _load_spend()
    data["_budget"] = float(value)
    _write_spend(data)


def _track_usage(n_in, n_out, model=None):
    u = st.session_state.setdefault("usage",
                                    {"in": 0, "out": 0, "calls": 0,
                                     "cost": 0.0})
    u["in"] += n_in
    u["out"] += n_out
    u["calls"] += 1
    p_in, p_out = PRICES.get(model, (0.0, 0.0))
    delta = n_in / 1e6 * p_in + n_out / 1e6 * p_out
    u["cost"] = u.get("cost", 0.0) + delta
    if st.session_state.get("backend") != "max":
        _add_spend(delta)


def _claude_cli_candidates():
    """Claude Code binaries the Agent SDK could run: the copy bundled inside
    claude-agent-sdk (which the SDK prefers by default) and any installed
    `claude` on PATH or in the usual locations."""
    import os
    import shutil
    from pathlib import Path
    cands = []
    try:
        import claude_agent_sdk
        b = (Path(claude_agent_sdk.__file__).parent / "_bundled"
             / ("claude.exe" if os.name == "nt" else "claude"))
        if b.is_file():
            cands.append(("bundled in claude-agent-sdk", str(b)))
    except Exception:
        pass
    seen = {c[1] for c in cands}
    home = Path.home()
    for p in [shutil.which("claude"), home / ".claude/local/claude",
              home / ".npm-global/bin/claude", "/usr/local/bin/claude",
              "/opt/homebrew/bin/claude", home / ".local/bin/claude"]:
        p = str(p) if p else ""
        if p and Path(p).is_file() and p not in seen:
            cands.append(("installed claude", p))
            seen.add(p)
    return cands


def _claude_cli_version(path):
    import re
    import subprocess
    try:
        out = subprocess.run([path, "--version"], capture_output=True, text=True,
                             timeout=20).stdout
        m = re.search(r"(\d+)\.(\d+)\.(\d+)", out or "")
        return m.group(0) if m else ""
    except Exception:
        return ""


def _ver_tuple(v):
    import re
    return tuple(int(x) for x in re.findall(r"\d+", str(v))[:3]) or (0,)


MAX_CLI_MIN_VERSION = "2.1.251"      # Fable 5.1 needs at least this Claude Code


def claude_code_status(refresh=False):
    """The newest Claude Code found (path, version, origin), cached for the
    session. Max mode runs through this binary, so a stale copy bundled in
    an old claude-agent-sdk no longer blocks new models when a newer
    `claude` is installed - and vice versa."""
    key = "_claude_code_status"
    if not refresh and key in st.session_state:
        return st.session_state[key]
    best = {}
    for origin, p in _claude_cli_candidates():
        v = _claude_cli_version(p)
        if v and (not best or _ver_tuple(v) > _ver_tuple(best["version"])):
            best = {"path": p, "version": v, "origin": origin}
    st.session_state[key] = best
    return best


def _max_update_command():
    import sys
    return f"{sys.executable} -m pip install -U claude-agent-sdk"


def _max_unsupported_hint(err, model):
    """Turn Claude Code's 'does not support this model; version X or newer
    is required' into a message that says exactly what to run."""
    import re
    m = re.search(r"does not support this model;?\s*version\s*([\d.]+)\s*or newer",
                  str(err), re.I)
    if not m:
        return None
    status = claude_code_status()
    have = status.get("version", "unknown")
    return (f"Claude Code {have} ({status.get('origin', 'not found')}) is too old "
            f"for {model}: version {m.group(1)} or newer is required. Fix in "
            f"Terminal, then restart the app:\n    {_max_update_command()}\n"
            "(claude-agent-sdk ships its own Claude Code; that command updates "
            "it. `claude update` works too - GrapheAI uses the newest Claude "
            "Code it finds.) Until then choose 'Fast (Haiku 4.5)' or switch the "
            "sidebar to 'API key'.")


def call_claude_max(system, user_msg, model, max_tokens=None):
    """Run one Claude call through the Claude Agent SDK (Claude Code), so it
    counts against the user's Claude Max subscription instead of API billing.

    Needs one-time setup (see MAX_SETUP.txt): `pip install claude-agent-sdk`
    (it bundles Claude Code) and a Claude Code login with the Max account.
    Personal use of your own subscription only. max_tokens is not enforced on
    this path (the CLI manages output length itself). New models need a
    recent Claude Code: the app runs the newest copy it can find.
    """
    try:
        from claude_agent_sdk import (query, ClaudeAgentOptions,
                                      AssistantMessage, TextBlock,
                                      ResultMessage)
    except ImportError as exc:
        raise RuntimeError(
            "Claude Max mode needs the Agent SDK. One-time setup (Terminal):\n"
            f"1) {_max_update_command()}\n"
            "2) Run `claude` once and log in with your Claude (Max) account\n"
            "Or switch the sidebar back to 'API key'. See MAX_SETUP.txt."
        ) from exc
    import asyncio
    status = claude_code_status()
    extra = {"cli_path": status["path"]} if status.get("path") else {}

    async def _run():
        opts = ClaudeAgentOptions(
            system_prompt=system,
            model=model,
            max_turns=1,
            # Pure text generation - keep Claude Code's tools out of the way.
            disallowed_tools=["Bash", "Edit", "Write", "Read", "Glob", "Grep",
                              "WebSearch", "WebFetch", "Task", "NotebookEdit"],
            **extra,
        )
        parts = []
        usage = None
        async for message in query(prompt=user_msg, options=opts):
            if isinstance(message, AssistantMessage):
                for block in message.content:
                    if isinstance(block, TextBlock):
                        parts.append(block.text)
            elif isinstance(message, ResultMessage):
                usage = getattr(message, "usage", None)
                if getattr(message, "is_error", False):
                    raise RuntimeError(getattr(message, "result", None)
                                       or "Claude Code returned an error.")
        return "".join(parts), usage

    try:
        text, usage = asyncio.run(_run())
    except Exception as exc:
        hint = _max_unsupported_hint(exc, model)
        if hint:
            raise RuntimeError(hint) from exc
        raise
    if not text.strip():
        raise RuntimeError(
            "Claude Code returned no text. If this mentions usage limits, "
            "your Max plan may be at its cap - switch the sidebar to "
            "'API key' to continue, or wait for the limit to reset.")
    try:
        _track_usage(int(usage.get("input_tokens", 0)),
                     int(usage.get("output_tokens", 0)))
    except Exception:
        _track_usage(0, 0)
    return text


# Models that reason before answering: their replies contain thinking blocks
# as well as text, and the thinking shares the max_tokens budget.
THINKING_MODELS = ("claude-fable-5", "claude-mythos-5", "claude-opus-5",
                   "claude-opus-4-8", "claude-opus-4-7", "claude-sonnet-5")


def _effective_max_tokens(model, requested):
    mt = requested or 1500
    if any(str(model).startswith(m) for m in THINKING_MODELS):
        mt = max(mt, 32000)
    return mt


def _response_text(resp):
    if getattr(resp, "stop_reason", None) == "refusal":
        raise RuntimeError(
            "Claude declined this request. Try rephrasing it, or choose a "
            "different model in the sidebar.")
    parts = [b.text for b in resp.content
             if getattr(b, "type", "") == "text" and getattr(b, "text", "")]
    if not parts:
        raise RuntimeError(
            "The model returned no text — it may have spent the whole token "
            "budget reasoning. Try a smaller model.")
    return "".join(parts)


EFFORT_LEVELS = ("low", "medium", "high", "xhigh", "max")
STREAM_ABOVE = 8000      # stream responses above this many output tokens


def _model_family(model):
    m = str(model)
    if m.startswith(("claude-fable-5", "claude-mythos-5")):
        return "fable"          # thinking always on; depth via effort
    if m.startswith("claude-opus-5"):
        return "opus5"          # adaptive by default; effort accepted
    if m.startswith(("claude-opus-4-8", "claude-opus-4-7", "claude-sonnet-5")):
        return "adaptive"       # must ask for adaptive thinking
    if m.startswith(("claude-opus-4-6", "claude-sonnet-4-6")):
        return "adaptive46"     # adaptive; effort up to 'high'/'max'
    return "classic"            # Haiku 4.5 and older: no thinking


def _api_kwargs(model, system, user_msg, max_tokens):
    """Request parameters tuned per model family. Fable 5.1: thinking is
    always on and its depth is set with output_config.effort (the
    sidebar's 'Reasoning effort'); Opus 5 the same; Opus 4.8 / Sonnet 5
    need adaptive thinking requested explicitly; Haiku gets none."""
    fam = _model_family(model)
    kw = {"model": model,
          "max_tokens": _effective_max_tokens(model, max_tokens),
          "system": system,
          "messages": [{"role": "user", "content": user_msg}]}
    effort = st.session_state.get("effort", "xhigh")
    if effort not in EFFORT_LEVELS:
        effort = "xhigh"
    if fam in ("fable", "opus5"):
        kw["output_config"] = {"effort": effort}
    elif fam == "adaptive":
        kw["thinking"] = {"type": "adaptive"}
        kw["output_config"] = {"effort": effort}
    elif fam == "adaptive46":
        kw["thinking"] = {"type": "adaptive"}
        kw["output_config"] = {"effort": "high" if effort == "xhigh"
                               else effort}
    return kw


def _api_create(client, kw):
    """Send the request with graceful degradation: refusal fallbacks
    (beta), then effort/adaptive thinking, then the bare request - so an
    older installed SDK or a model that rejects a parameter still works.
    Long outputs are streamed to avoid HTTP timeouts."""
    import anthropic
    stream = kw["max_tokens"] > STREAM_ABOVE
    fam = _model_family(kw["model"])
    attempts = []
    if fam in ("fable", "opus5") and not st.session_state.get("_no_fallbacks"):
        attempts.append(("beta", dict(kw, betas=["server-side-fallback-2026-07-01"],
                                      fallbacks="default")))
    attempts.append(("plain", dict(kw)))
    attempts.append(("minimal", {k: kw[k] for k in
                                 ("model", "max_tokens", "system", "messages")}))
    last = None
    for kind, params in attempts:
        try:
            api = client.beta.messages if kind == "beta" else client.messages
            if stream:
                with api.stream(**params) as s:
                    return s.get_final_message()
            return api.create(**params)
        except TypeError as e:                 # installed SDK lacks a parameter
            last = e
            if kind == "beta":
                st.session_state["_no_fallbacks"] = True
            continue
        except anthropic.BadRequestError as e:
            last = e
            low = str(e).lower()
            if kind == "beta":
                st.session_state["_no_fallbacks"] = True
                continue
            if kind == "plain" and any(w in low for w in
                                       ("output_config", "effort", "thinking",
                                        "fallback", "beta", "unexpected")):
                continue
            raise
    raise last


def call_claude(api_key, system, user_msg, model, max_tokens=None):
    """Single Claude call via the selected backend; tracks token usage.
    Backend 'api': the Anthropic API (Fable 5.1 with effort-controlled
    reasoning, streaming, refusal fallbacks). Backend 'max': the Claude
    Agent SDK, billed to the Claude Max plan."""
    if st.session_state.get("backend") == "max":
        return call_claude_max(system, user_msg, model, max_tokens)
    import anthropic
    client = anthropic.Anthropic(api_key=api_key)
    resp = _api_create(client, _api_kwargs(model, system, user_msg,
                                           max_tokens))
    _track_usage(resp.usage.input_tokens, resp.usage.output_tokens, model)
    return _response_text(resp)


# --------------------------------------------------------------------------
# News store + signal extraction
# --------------------------------------------------------------------------
NEWS_FILE = ANSWERS_DIR / "pv_radar.json"
NEWS_CATEGORIES = ["Research", "Technology", "Funding", "Market", "Policy",
                   "Manufacturing", "Patent"]
NEWS_SYSTEM = """\
You extract structured signals from ONE news item about solar photovoltaics
(especially perovskite). You get a headline and a short summary. Respond with
ONLY a JSON object (no markdown), keys exactly:
{"efficiency_pct": number|null, "cell_type": "single junction"|"tandem"|
"module"|"unspecified", "funding_musd": number|null, "company": string|null,
"country": string|null, "category": one of Research|Technology|Funding|Market|
Policy|Manufacturing}

Rules: efficiency_pct only if a solar-cell or module power-conversion
efficiency percentage is actually stated (else null). funding_musd = any
investment/funding/grant/raise amount converted to US$ millions
($6 million -> 6, "€1.2 billion" -> about 1300); else null. company = the main
organisation if clearly named, else null. country if clearly implied, else
null. Pick the single best category. Never invent numbers that are not in the
text."""


def fetch_google_news(query, n=40):
    """Recent news items via Google News RSS (no API key)."""
    import requests as _rq
    import email.utils as _eu
    import xml.etree.ElementTree as ET
    from urllib.parse import quote
    url = ("https://news.google.com/rss/search?q=" + quote(query)
           + "&hl=en-US&gl=US&ceid=US:en")
    r = _rq.get(url, timeout=30,
                headers={"User-Agent": "Mozilla/5.0 (PV Radar)"})
    r.raise_for_status()
    root = ET.fromstring(r.content)
    items = []
    for it in root.iter("item"):
        title = (it.findtext("title") or "").strip()
        if not title:
            continue
        desc = it.findtext("description") or ""
        summary = re.sub(r"\s+", " ", re.sub("<[^>]+>", " ", desc)).strip()[:600]
        src_el = it.find("source")
        source = (src_el.text.strip() if src_el is not None and src_el.text
                  else "")
        pub = it.findtext("pubDate") or ""
        try:
            date = _eu.parsedate_to_datetime(pub).date().isoformat()
        except Exception:
            date = ""
        items.append({"title": title, "summary": summary,
                      "link": (it.findtext("link") or "").strip(),
                      "source": source, "date": date})
        if len(items) >= n:
            break
    return items


def load_news():
    import json as _json
    if NEWS_FILE.exists():
        try:
            return _json.loads(NEWS_FILE.read_text(encoding="utf-8"))
        except Exception:
            return []
    return []


def save_news(rows):
    import json as _json
    try:
        ANSWERS_DIR.mkdir(exist_ok=True)
        NEWS_FILE.write_text(_json.dumps(rows, ensure_ascii=False),
                             encoding="utf-8")
    except Exception:
        pass


# Watchlists (saved searches), keyword alerts and RSS feeds - one file.
WATCH_FILE = ANSWERS_DIR / "pv_watchlists.json"

# Dedicated PV outlets, pulled alongside Google News. Editable in the app.
DEFAULT_FEEDS = ["https://www.pv-magazine.com/feed/",
                 "https://www.pv-tech.org/feed/"]


def load_watch():
    import json as _json
    if WATCH_FILE.exists():
        try:
            d = _json.loads(WATCH_FILE.read_text(encoding="utf-8"))
            return {"queries": list(d.get("queries", [])),
                    "alerts": list(d.get("alerts", [])),
                    "patents": list(d.get("patents", [])),
                    "feeds": (list(d["feeds"]) if "feeds" in d
                              else list(DEFAULT_FEEDS))}
        except Exception:
            pass
    return {"queries": [], "alerts": [], "patents": [],
            "feeds": list(DEFAULT_FEEDS)}


def save_watch(d):
    import json as _json
    try:
        ANSWERS_DIR.mkdir(exist_ok=True)
        WATCH_FILE.write_text(_json.dumps(d, ensure_ascii=False),
                              encoding="utf-8")
    except Exception:
        pass


def fetch_rss(url, n=30):
    """Items from one dedicated RSS feed (PV Magazine, PV-Tech, ...)."""
    import requests as _rq
    import email.utils as _eu
    import xml.etree.ElementTree as ET
    r = _rq.get(url, timeout=30,
                headers={"User-Agent": "Mozilla/5.0 (PV Radar)"})
    r.raise_for_status()
    root = ET.fromstring(r.content)
    feed_title = (root.findtext("channel/title") or "").strip()
    items = []
    for it in root.iter("item"):
        title = (it.findtext("title") or "").strip()
        if not title:
            continue
        desc = it.findtext("description") or ""
        summary = re.sub(r"\s+", " ",
                         re.sub("<[^>]+>", " ", desc)).strip()[:600]
        pub = it.findtext("pubDate") or ""
        try:
            date = _eu.parsedate_to_datetime(pub).date().isoformat()
        except Exception:
            date = ""
        items.append({"title": title, "summary": summary,
                      "link": (it.findtext("link") or "").strip(),
                      "source": feed_title or url, "date": date})
        if len(items) >= n:
            break
    return items


def analyse_new(items, extract_model, api_key, progress_cb=None, tag="",
                force_category=None):
    """Analyse items not yet in the store and append them.
    Returns (added, store_size)."""
    store = load_news()
    seen = {r.get("link") or r.get("title") for r in store}
    todo = [it for it in items
            if (it.get("link") or it["title"]) not in seen]
    for i, it in enumerate(todo, start=1):
        if progress_cb:
            progress_cb(i, len(todo), it["title"])
        try:
            sig = extract_news_signals(it, api_key, extract_model)
        except Exception:
            sig = {"efficiency_pct": None, "cell_type": "unspecified",
                   "funding_musd": None, "company": None,
                   "country": None, "category": "Research"}
        if force_category:
            sig["category"] = force_category
        store.append({**it, **sig, "query": tag})
    if todo:
        save_news(store)
    return len(todo), len(store)


# --------------------------------------------------------------------------
# Startup benchmark: founder-entered facts per product segment
# --------------------------------------------------------------------------
BENCH_FILE = ANSWERS_DIR / "pv_benchmark.json"
BENCH_SEGMENTS = ["Single junction", "4T pero-Si tandem",
                  "2T monolithic tandem"]
BENCH_COLS = ["Company", "PV performance", "Funding (US$M)",
              "Product sales", "IP portfolio", "MoU/LoI signed",
              "Market entry barrier",
              "News: best eff %", "News: funding US$M"]


def _bench_row(company=""):
    r = {c: "" for c in BENCH_COLS}
    r["Company"] = company
    return r


def load_bench():
    import json as _json
    if BENCH_FILE.exists():
        try:
            d = _json.loads(BENCH_FILE.read_text(encoding="utf-8"))
            return {s: list(d.get(s, [])) for s in BENCH_SEGMENTS}
        except Exception:
            pass
    return {s: [_bench_row("My startup")] for s in BENCH_SEGMENTS}


def save_bench(d):
    import json as _json
    try:
        ANSWERS_DIR.mkdir(exist_ok=True)
        BENCH_FILE.write_text(_json.dumps(d, ensure_ascii=False),
                              encoding="utf-8")
    except Exception:
        pass


BENCHMARK_SYSTEM = """\
You benchmark a founder's PV startup against its competitors, per product
segment, for investor preparation. Inputs: per-segment tables the founder
maintains (their own venture is the first row, named "My startup" or
similar) plus recent analysed news items for context.

For EACH segment that has rows:
1. **Positioning matrix** - a condensed markdown table, companies as
   rows, the parameters as columns; write "unknown" for empty cells.
2. **Where the startup stands** - parameter by parameter: ahead /
   behind / unknown, quoting the decisive numbers.
3. **Moat & gaps** - what the data says the startup uniquely has, and
   the 3 most important gaps to close in that segment.

Then, once, across segments:
4. **Research checklist** - every "unknown" cell, with where to find the
   fact (patent registers such as Espacenet/Google Patents for IP;
   annual reports and press releases for sales; company statements for
   MoU/LoI; certification charts for efficiencies).
5. **Investor Q&A prep** - the 5 hardest questions this comparison
   invites, each with the honest answer the current data supports.

Rules - the founder's credibility depends on them:
- Use ONLY the tables and news items. NEVER invent or estimate a missing
  value; an empty cell is "unknown", said loudly.
- Distinguish the founder's targets from achieved values wherever the
  table marks them.
- News-derived efficiencies are company-reported unless stated
  certified; say so.
- No investment-advice language."""


def run_fetch(query, n, extract_model, api_key, progress_cb=None,
              freshness=""):
    """Fetch one Google News search and analyse the new items.
    freshness is a Google News window operator like 'when:7d' ('' = any)."""
    q = f"{query} {freshness}".strip()
    return analyse_new(fetch_google_news(q, n), extract_model,
                       api_key, progress_cb, tag=query)


def run_feed(url, n, extract_model, api_key, progress_cb=None):
    """Fetch one RSS feed and analyse the new items."""
    return analyse_new(fetch_rss(url, n), extract_model,
                       api_key, progress_cb, tag=f"feed:{url}")


def patent_url(query):
    """WIPO Patentscope RSS for a full-text patent search."""
    from urllib.parse import quote
    return ("https://patentscope.wipo.int/search/rss.jsf?query="
            + quote(f'EN_ALLTXT:("{query}")')
            + "&sortOption=Pub+Date+Desc")


def run_patents(query, n, extract_model, api_key, progress_cb=None):
    """Fetch one Patentscope search and analyse the new filings;
    stored with category 'Patent'."""
    return analyse_new(fetch_rss(patent_url(query), n), extract_model,
                       api_key, progress_cb, tag=f"patent:{query}",
                       force_category="Patent")


def extract_news_signals(item, api_key, model):
    import json as _json
    raw = call_claude(api_key, NEWS_SYSTEM,
                      f"HEADLINE: {item['title']}\nSUMMARY: {item['summary']}",
                      model, max_tokens=400)
    raw = re.sub(r"^```(json)?|```$", "", raw.strip(), flags=re.M).strip()
    data = _json.loads(raw)
    cat = data.get("category")
    if cat not in NEWS_CATEGORIES:
        cat = "Research"
    return {
        "efficiency_pct": data.get("efficiency_pct"),
        "cell_type": data.get("cell_type") or "unspecified",
        "funding_musd": data.get("funding_musd"),
        "company": data.get("company"),
        "country": data.get("country"),
        "category": cat,
    }


COMPANY_REPORT_SYSTEM = """\
You write a short competitive-intelligence briefing on ONE organisation's
photovoltaic / perovskite activity, from a list of that organisation's news
items (each: date, headline, category, and any efficiency or funding figure).

Produce:
1. A 4-6 sentence summary: technology focus, key milestones, funding,
   partnerships, and momentum.
2. "Notable items" - up to 8 bullets, each "YYYY-MM-DD - one line".
Use ONLY the provided items; never invent figures. Call out gaps explicitly
(e.g. "no efficiency values reported in these items")."""

DIGEST_SYSTEM = """\
You write a periodic intelligence digest on photovoltaics / perovskite
news for a senior PV researcher, from a structured list of analysed news
items (date | category | company | country | efficiency % | funding US$M
| headline).

Structure (markdown):
# PV Radar digest - {period}
1. **The week at a glance** - 3-5 sentences on the dominant themes.
2. **Efficiency & technology** - notable results; state the numbers,
   cell types, and who reported them.
3. **Funding & industry** - amounts, companies, countries; sum notable
   totals.
4. **Policy & market** - only if such items exist.
5. **Watch next** - 2-3 bullets on what these items suggest is coming.

Rules: use ONLY the provided items; never invent numbers or events; name
sources by company/headline so items can be found again; if a section has
no items, drop it. Keep it under ~500 words."""

# Approximate country centroids (lat, lon) for the map. Extend as needed.
COUNTRY_CENTROIDS = {
    "united states": (39.8, -98.6), "usa": (39.8, -98.6),
    "us": (39.8, -98.6), "china": (35.9, 104.2),
    "south korea": (36.5, 127.9), "korea": (36.5, 127.9),
    "japan": (36.2, 138.3), "india": (22.4, 78.9),
    "germany": (51.2, 10.4), "united kingdom": (54.0, -2.0),
    "uk": (54.0, -2.0), "france": (46.6, 2.2),
    "switzerland": (46.8, 8.2), "netherlands": (52.1, 5.3),
    "belgium": (50.6, 4.6), "spain": (40.2, -3.7),
    "italy": (42.8, 12.6), "sweden": (62.2, 15.6),
    "australia": (-25.7, 133.8), "saudi arabia": (23.9, 45.1),
    "uae": (23.4, 53.8), "united arab emirates": (23.4, 53.8),
    "singapore": (1.35, 103.8), "canada": (56.1, -106.3),
    "poland": (51.9, 19.1), "austria": (47.6, 14.1),
    "denmark": (56.0, 9.5), "norway": (60.5, 8.5),
    "finland": (61.9, 25.7), "israel": (31.0, 34.9),
    "taiwan": (23.7, 121.0), "brazil": (-14.2, -51.9),
    "ireland": (53.4, -8.2), "portugal": (39.4, -8.2),
    "czech republic": (49.8, 15.5), "greece": (39.1, 21.8),
    "turkey": (38.9, 35.2), "russia": (61.5, 105.3),
    "mexico": (23.6, -102.6), "south africa": (-30.6, 22.9),
    "new zealand": (-41.8, 172.7), "hong kong": (22.3, 114.2),
    "qatar": (25.3, 51.2), "egypt": (26.8, 30.8),
    "morocco": (31.8, -7.1), "chile": (-35.7, -71.5),
}


def country_latlon(name):
    if not name:
        return None
    return COUNTRY_CENTROIDS.get(str(name).strip().lower())


def _md_runs(par, text):
    """Render **bold**, *italic* and `code` inline markup as Word runs."""
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`", text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()])
        if m.group(1) is not None:
            par.add_run(m.group(1)).bold = True
        elif m.group(2) is not None:
            par.add_run(m.group(2)).italic = True
        else:
            r = par.add_run(m.group(3))
            r.font.name = "Consolas"
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:])


def md_to_docx(doc, md):
    """Append markdown to a Document as real Word formatting: headings,
    bold/italic, lists, and tables - instead of raw **, # and | text."""
    from docx.shared import Pt
    lines = md.splitlines()
    i, in_code = 0, False
    while i < len(lines):
        stripped = lines[i].strip()
        if stripped.startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            r = doc.add_paragraph().add_run(lines[i])
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            i += 1
            continue
        if not stripped or re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            i += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            tbl = []
            while i < len(lines):
                s = lines[i].strip()
                if not (s.startswith("|") and s.endswith("|")):
                    break
                cells = [c.strip() for c in s.strip("|").split("|")]
                if not all(set(c) <= set("-: ") for c in cells):
                    tbl.append(cells)
                i += 1
            if tbl:
                ncols = max(len(r) for r in tbl)
                t = doc.add_table(rows=len(tbl), cols=ncols)
                try:
                    t.style = "Table Grid"
                except Exception:
                    pass
                for ri, row in enumerate(tbl):
                    for ci in range(ncols):
                        cpar = t.cell(ri, ci).paragraphs[0]
                        _md_runs(cpar, row[ci] if ci < len(row) else "")
                        if ri == 0:
                            for r_ in cpar.runs:
                                r_.bold = True
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if m:
            h = doc.add_heading("", level=min(len(m.group(1)), 4))
            _md_runs(h, m.group(2).strip())
            i += 1
            continue
        m = re.match(r"^[-*+]\s+(.*)", stripped)
        if m:
            _md_runs(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue
        m = re.match(r"^\d+[.)]\s+(.*)", stripped)
        if m:
            _md_runs(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue
        _md_runs(doc.add_paragraph(), stripped)
        i += 1


def save_report(name, text):
    """Save a briefing to the shared answers folder as .md and .docx."""
    ANSWERS_DIR.mkdir(exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    safe = re.sub(r"[^A-Za-z0-9 _-]", "", name)[:60].strip() or "report"
    (ANSWERS_DIR / f"PVRadar_{safe}_{stamp}.md").write_text(
        f"# PV Radar briefing: {name}\n\n{text}\n", encoding="utf-8")
    try:
        from docx import Document
        doc = Document()
        doc.add_heading(f"PV Radar briefing: {name}", level=1)
        md_to_docx(doc, text)
        doc.save(ANSWERS_DIR / f"PVRadar_{safe}_{stamp}.docx")
    except Exception:
        pass


# --------------------------------------------------------------------------
# Page, theme, sidebar
# --------------------------------------------------------------------------
st.set_page_config(page_title="PV Radar - by GrapheAI",
                   page_icon="📡", layout="wide")

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Serif:wght@600&family=IBM+Plex+Sans:wght@400;500;600&family=IBM+Plex+Mono:wght@400;500&display=swap');

/* ---- Dark technical: control-room slate, signal-orange accent ---- */
html, body, [class*="css"], .stMarkdown, p, li {
    font-family: 'IBM Plex Sans', 'Segoe UI', sans-serif;
    font-size: 16.5px;
    color: #E6EAF0;
}
.stApp { background: #12161C; }

/* Header: radar console masthead */
.pv-header { border-bottom: 3px solid #FF6B3D; padding-bottom: 12px;
             margin-bottom: 10px; }
.pv-title  { font-family: 'IBM Plex Serif', Georgia, serif;
             font-size: 2.6rem; font-weight: 600; color: #F4F6F9;
             margin: 0; letter-spacing: -0.5px; }
.pv-sub    { font-size: 0.82rem; color: #8E99A8; margin-top: 6px;
             text-transform: uppercase; letter-spacing: 2px; }
.pv-sub b  { color: #FF8A5C; font-weight: 600; }

/* Tabs */
.stTabs [data-baseweb="tab-list"] { border-bottom: 1px solid #242D38;
                                    gap: 2px; }
.stTabs [data-baseweb="tab"] {
    font-size: 1.05rem; font-weight: 500; padding: 12px 18px;
    color: #8E99A8;
}
.stTabs [data-baseweb="tab"]:hover { color: #FFB08F; }
.stTabs [aria-selected="true"] { color: #FF8A5C; font-weight: 600; }
.stTabs [data-baseweb="tab-highlight"] { background-color: #FF6B3D;
                                         height: 3px; }

/* Sidebar */
section[data-testid="stSidebar"] { background: #171D25;
    border-right: 1px solid #242D38; }
section[data-testid="stSidebar"] * { font-size: 15px; }
section[data-testid="stSidebar"] h1 { font-family: 'IBM Plex Serif', serif;
                                      font-size: 1.4rem; color: #F4F6F9; }

/* Headings */
h2 { font-family: 'IBM Plex Serif', Georgia, serif; font-weight: 600;
     color: #DFE5EC; }
h3 { color: #C3CBD6; }

/* Panels: forms, expanders, dataframes */
[data-testid="stForm"] {
    background: #1A212B; border: 1px solid #263140; border-radius: 14px;
    padding: 1.1rem 1.3rem .9rem; box-shadow: 0 2px 8px rgba(0,0,0,.35);
}
details { border: 1px solid #263140; border-radius: 12px;
          background: #1A212B; }
[data-testid="stExpander"] summary { font-size: 1.02rem; font-weight: 500; }
[data-testid="stDataFrame"] { border: 1px solid #263140;
                              border-radius: 12px; }

/* Metrics: instrument readouts */
[data-testid="stMetric"] {
    background: #1A212B; border: 1px solid #263140; border-radius: 12px;
    padding: .7rem .95rem; box-shadow: 0 2px 6px rgba(0,0,0,.3);
}
[data-testid="stMetricValue"] { font-family: 'IBM Plex Mono', monospace;
                                color: #FF8A5C; }
[data-testid="stMetricLabel"] { color: #8E99A8;
                                text-transform: uppercase;
                                letter-spacing: 1px; font-size: .78rem; }

/* Buttons */
.stButton>button, .stDownloadButton>button, .stFormSubmitButton>button {
    font-size: 1.0rem; font-weight: 600; border-radius: 10px;
    padding: 0.5rem 1.25rem;
}
.stButton>button:hover, .stFormSubmitButton>button:hover {
    filter: brightness(1.12);
}

/* Inputs */
.stTextArea textarea, .stTextInput input { font-size: 1.02rem;
                                           border-radius: 10px; }
</style>
""", unsafe_allow_html=True)

st.markdown(
    "<div class='pv-header'>"
    "<p class='pv-title'>📡 PV Radar</p>"
    "<p class='pv-sub'>by <b>GrapheAI</b> · developed by "
    "<b>Dr. Anurag Krishna</b></p>"
    "</div>",
    unsafe_allow_html=True)

with st.sidebar:
    st.title("📡 PV Radar")
    st.caption("by GrapheAI · Dr. Anurag Krishna")
    backend_label = st.radio(
        "Claude access", ["API key (pay per use)",
                          "Claude Max subscription (needs Claude Code)"])
    st.session_state["backend"] = ("max" if "Max" in backend_label else "api")
    if st.session_state["backend"] == "api":
        api_key = st.text_input("Anthropic API key", type="password")
    else:
        # Sentinel so features unlock; never sent anywhere in Max mode.
        api_key = "claude-max-subscription"
        _cc = claude_code_status()
        if not _cc:
            st.warning("Claude Code not found for Max mode - run "
                       f"`{_max_update_command()}` and log in with `claude`.")
        elif _ver_tuple(_cc["version"]) < _ver_tuple(MAX_CLI_MIN_VERSION):
            st.warning(f"Claude Code {_cc['version']} ({_cc['origin']}) is older "
                       f"than {MAX_CLI_MIN_VERSION}, which Fable 5.1 needs. Run in "
                       f"Terminal, then restart: `{_max_update_command()}`. Until "
                       "then choose Haiku 4.5 or the API-key backend.")
        else:
            st.caption(f"Claude Code {_cc['version']} ({_cc['origin']})")
    st.session_state["effort"] = st.select_slider(
        "Reasoning effort", options=list(EFFORT_LEVELS),
        value=st.session_state.get("effort", "xhigh"),
        help="Thinking depth for Fable 5.1 / Opus 5 / Sonnet 5 on the API "
             "backend. xhigh = professor-grade default; max = when "
             "correctness matters more than time; low = quick lookups.")
    model_label = st.selectbox("Model (briefings)", list(MODELS.keys()),
                               index=1)
    model = MODELS[model_label]
    do_autosave = st.toggle("Auto-save briefings (.docx + .md)", value=True)
    st.markdown("---")
    st.caption(f"News store: {len(load_news())} items "
               f"(shared with GrapheAI: answers/pv_radar.json)")

    u = st.session_state.get("usage")
    if u:
        line = (f"Session: {u['calls']} calls - "
                f"{u['in']:,} in / {u['out']:,} out tokens")
        if st.session_state.get("backend") != "max" and u.get("cost"):
            line += f"  ·  ≈ ${u['cost']:.2f}"
        st.caption(line)

    # Monthly spend + budget cap, shared with GrapheAI (same spend.json).
    if st.session_state.get("backend") != "max":
        ms = month_spend()
        saved_budget = get_budget()
        budget = st.number_input(
            "Monthly budget $ (0 = off)", min_value=0.0,
            value=float(saved_budget), step=5.0, key="budget_input",
            help="Shared with GrapheAI - one budget across both apps.")
        if budget != saved_budget:
            set_budget(budget)
        if budget > 0:
            frac = ms / budget if budget else 0
            st.progress(min(frac, 1.0),
                        text=f"This month: ${ms:.2f} / ${budget:.0f}")
            if ms >= budget:
                st.error("Over your monthly budget.")
            elif frac >= 0.8:
                st.warning("Approaching your monthly budget.")
        else:
            st.caption(f"This month (API): ≈ ${ms:.2f}")

# --------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------
REPORT_SYSTEM = """\
You write a detailed competitive/technology intelligence report on
photovoltaics / perovskite for a senior PV researcher, from a structured
list of analysed news items (date | category | company | country |
efficiency % | funding US$M | headline).

Write a thorough markdown report with these sections (drop a section only
if no items support it):
# {title}
1. **Executive summary** - 5-8 sentences.
2. **Technology landscape** - efficiency results and technical milestones,
   grouped by cell type; state the numbers, who reported them, and when.
3. **Key players** - the most active organisations and what each is doing.
4. **Funding & investment** - amounts, recipients, totals worth noting.
5. **Geography** - where the activity concentrates.
6. **Policy & market context** - if such items exist.
7. **Assessment & outlook** - your synthesis: momentum, gaps, and what to
   watch next.
8. **Key items** - up to 15 bullets, each "YYYY-MM-DD - one line (company)".

Rules: use ONLY the provided items - never invent events, numbers, or
names; where the evidence is thin, say so explicitly; aim for about
{words} words of substance with no filler."""

INVESTOR_TEMPLATES = {
    "Market report": """\
1. **Executive summary** - the market story in 6-8 sentences.
2. **Market signals & momentum** - activity volume over the period,
   funding totals, geographic spread, which segments are heating up.
3. **Commercialization status** - who is at lab / pilot / production;
   scale-up, factory and product announcements with dates.
4. **Key players by segment** - grouped (tandem, single junction,
   modules, equipment/materials), one line each on what they just did.
5. **Risks & open questions** - what the news says is still unsolved.
6. **Outlook** - where the momentum points over the next 12-24 months.
7. **Key sources** - up to 15 dated bullets.""",

    "Competitive analysis": """\
1. **Executive summary** - the competitive picture in 5-8 sentences.
2. **Competitor profiles** - for each significant company: technology
   approach, best reported efficiency, funding raised, notable
   milestones with dates, and momentum (active / quiet lately).
3. **Positioning table**:
   | Company | Technology | Best reported eff. | Funding US$M | Stage | Recent momentum |
4. **White space** - segments or approaches with conspicuously little
   activity in these items.
5. **Implications for a new entrant** - where the openings are, judged
   strictly from the coverage above.
6. **Key sources** - up to 15 dated bullets.""",

    "Investment & funding landscape": """\
1. **Executive summary** - the funding climate in 5-8 sentences.
2. **Funding rounds table**:
   | Date | Company | Amount US$M | Country | What the money is for |
3. **Totals & trends** - by geography and over time within the period.
4. **Who is investing** - investors/agencies where the items name them.
5. **Signals for fundraising** - what kinds of stories and stages got
   funded, typical amounts.
6. **Key sources** - up to 15 dated bullets.""",

    "Technology landscape": """\
1. **Executive summary** - the technical state of play in 5-8 sentences.
2. **Efficiency state of play** - best reported values by cell type,
   who and when; flag which are certified vs company-reported.
3. **Manufacturing & scale-up** - deposition routes, line capacities,
   pilot production in the news.
4. **Stability & durability signals** - what the items report.
5. **Technology risks** - unresolved issues visible in the coverage.
6. **Direction of travel** - where the field is heading.
7. **Key sources** - up to 15 dated bullets.""",
}

INVESTOR_SYSTEM = """\
You write an investor-facing intelligence report on the perovskite / PV
sector, for a researcher preparing startup materials. Source: a
structured list of analysed news items (date | category | company |
country | efficiency % | funding US$M | headline).

AUDIENCE: {audience}

Write a polished, decision-oriented markdown report titled "{title}",
about {words} words, with exactly this structure:
{structure}

Rules - these protect the author's credibility in front of investors:
- Use ONLY the provided items. Never invent companies, amounts, rounds,
  efficiencies, or events. Every figure must be traceable to an item
  (name the company and date inline).
- Distinguish company-claimed values from certified/verified ones
  whenever the wording allows; if unclear, say "company-reported".
- Where coverage is thin for a section, say so plainly - investors
  punish overclaiming far more than gaps.
- End with a short **Data & limitations** note: this is news-derived
  intelligence over the stated period, not comprehensive market
  research; figures need primary-source verification before use in a
  pitch deck.
- No investment-advice language."""

st.caption("News via Google News (no key needed) · Claude extracts "
           "efficiency, funding, company, country and category · "
           "everything accumulates locally in answers/pv_radar.json")

rows = load_news()
import pandas as _pd
try:
    import altair as alt
except Exception:
    alt = None

view = None
if rows:
    df = _pd.DataFrame(rows)
    for col in ("title", "company", "country", "source", "category",
                "cell_type", "efficiency_pct", "funding_musd", "date",
                "link", "summary"):
        if col not in df.columns:
            df[col] = None
    for col in ("efficiency_pct", "funding_musd"):
        df[col] = _pd.to_numeric(df[col], errors="coerce")
    df["date_dt"] = _pd.to_datetime(df["date"], errors="coerce")

    with st.expander("🔎 Filter the data (applies to every tab)"):
        fc1, fc2, fc3, fc4 = st.columns([2, 2, 2, 1])
        with fc1:
            fsearch = st.text_input("Find (title/company/country)",
                                    key="radar_filter")
        with fc2:
            fcats = st.multiselect("Category", NEWS_CATEGORIES,
                                   key="radar_cats")
        with fc3:
            countries = sorted([c for c in df["country"].dropna().unique()
                                if str(c).strip()])
            fcountry = st.multiselect("Country", countries,
                                      key="radar_country")
        with fc4:
            fperiod = st.selectbox("Period",
                                   ["Any time", "30 days", "90 days",
                                    "12 months"], key="radar_period")
    view = df.copy()
    if fsearch.strip():
        s = fsearch.strip().lower()
        hay = (view["title"].fillna("").astype(str) + " "
               + view["company"].fillna("").astype(str) + " "
               + view["country"].fillna("").astype(str)).str.lower()
        view = view[hay.str.contains(s, regex=False)]
    if fcats:
        view = view[view["category"].isin(fcats)]
    if fcountry:
        view = view[view["country"].isin(fcountry)]
    if fperiod != "Any time" and view["date_dt"].notna().any():
        days = {"30 days": 30, "90 days": 90, "12 months": 365}[fperiod]
        cutoff = view["date_dt"].max() - _pd.Timedelta(days=days)
        view = view[view["date_dt"] >= cutoff]

(tab_dash, tab_fetch, tab_trend, tab_rec, tab_map, tab_comp,
 tab_rep, tab_bench, tab_arch) = st.tabs(
    ["📊 Dashboard", "📡 Fetch & Watchlists", "📈 Trends", "🏁 Records",
     "🗺️ Map", "🏢 Companies", "📑 Reports", "🥊 Benchmark",
     "🗂️ Archive"])

NO_DATA = "No radar data yet - run a fetch in **📡 Fetch & Watchlists**."

# ------------------------------ DASHBOARD ---------------------------------
with tab_dash:
    if view is None:
        st.info(NO_DATA)
    else:
        # Keyword alerts: flag matching items from the last 30 days.
        alert_terms = load_watch()["alerts"]
        if alert_terms:
            hay = (df["title"].fillna("").astype(str) + " "
                   + df["company"].fillna("").astype(str) + " "
                   + df["summary"].fillna("").astype(str)).str.lower()
            mask = False
            for t in alert_terms:
                mask = mask | hay.str.contains(t.lower(), regex=False)
            recent = df["date_dt"].notna() & (
                df["date_dt"] >= df["date_dt"].max()
                - _pd.Timedelta(days=30))
            hits_a = df[mask & recent].sort_values("date_dt",
                                                   ascending=False)
            if not hits_a.empty:
                with st.expander(f"🚨 {len(hits_a)} alert match(es) in "
                                 "the last 30 days", expanded=True):
                    for _, r in hits_a.head(12).iterrows():
                        bits = [str(r.get("date") or "")]
                        if r.get("company"):
                            bits.append(str(r["company"]))
                        line = " · ".join(b for b in bits if b)
                        link = r.get("link") or ""
                        title = str(r.get("title") or "")
                        if link:
                            st.markdown(f"- **{line}** — "
                                        f"[{title}]({link})")
                        else:
                            st.markdown(f"- **{line}** — {title}")

        eff = view["efficiency_pct"].dropna()
        m1, m2, m3, m4, m5 = st.columns(5)
        m1.metric("Items", len(view))
        m2.metric("Companies", view["company"].dropna().nunique())
        m3.metric("Reporting eff.", int(eff.count()))
        m4.metric("Best eff.",
                  f"{eff.max():.1f}%" if not eff.empty else "-")
        fund_sum = view["funding_musd"].dropna().sum()
        m5.metric("Funding",
                  f"US${fund_sum/1000:.2f}B" if fund_sum >= 1000
                  else (f"US${fund_sum:,.0f}M" if fund_sum else "-"))

        xa, ya = st.columns(2)
        with xa:
            x_field = st.selectbox("X axis",
                                   ["Date", "Category", "Country",
                                    "Cell type"], key="radar_x")
        with ya:
            y_field = st.selectbox("Y axis",
                                   ["Efficiency (%)", "Funding (US$M)"],
                                   key="radar_y")
        ycol = "efficiency_pct" if "Effic" in y_field else "funding_musd"
        xmap = {"Date": "date_dt", "Category": "category",
                "Country": "country", "Cell type": "cell_type"}
        xcol = xmap[x_field]
        pdata = view.dropna(subset=[ycol])
        if alt is not None and not pdata.empty:
            xenc = (alt.X("date_dt:T", title=x_field)
                    if xcol == "date_dt"
                    else alt.X(f"{xcol}:N", title=x_field, sort="-y"))
            chart = alt.Chart(pdata).mark_point(
                size=90, filled=True, opacity=0.75).encode(
                x=xenc, y=alt.Y(f"{ycol}:Q", title=y_field),
                color=alt.Color("cell_type:N", title="Cell type"),
                tooltip=["title:N", "company:N", "country:N",
                         "efficiency_pct:Q", "funding_musd:Q", "date:N"])
            st.altair_chart(chart, use_container_width=True)
        else:
            st.caption("No data with that Y value in this view.")
        cc1, cc2 = st.columns(2)
        with cc1:
            st.markdown("**By category**")
            cat_ct = (view["category"].value_counts()
                      .rename_axis("category").reset_index(name="items"))
            if alt is not None:
                st.altair_chart(alt.Chart(cat_ct).mark_bar().encode(
                    x="items:Q", y=alt.Y("category:N", sort="-x"),
                    tooltip=["category", "items"]),
                    use_container_width=True)
            else:
                st.bar_chart(cat_ct.set_index("category"))
        with cc2:
            st.markdown("**Funding over time (US$M)**")
            ft = view.dropna(subset=["funding_musd", "date_dt"])
            if not ft.empty:
                ft = (ft.groupby(ft["date_dt"].dt.date)["funding_musd"]
                      .sum().reset_index())
                ft.columns = ["date", "funding"]
            if not ft.empty and alt is not None:
                st.altair_chart(alt.Chart(ft).mark_bar(
                    color="#FF7043").encode(
                    x="date:T", y=alt.Y("funding:Q", title="US$M"),
                    tooltip=["date", "funding"]),
                    use_container_width=True)
            elif not ft.empty:
                st.bar_chart(ft.set_index("date"))
            else:
                st.caption("No funding amounts in this view.")

# ------------------------------ FETCH -------------------------------------
with tab_fetch:
    st.markdown("**Fetch news** once or refresh every saved watchlist in "
                "one click. New items are analysed by Claude and added to "
                "the permanent store; duplicates are skipped automatically.")
    rc1, rc2, rc3, rc4 = st.columns([3, 1, 1, 1])
    with rc1:
        rquery = st.text_input("News search",
                               value="perovskite solar cell", key="radar_q",
                               help="Plain Google News keywords, e.g. "
                                    "'Renshine perovskite' or 'tandem "
                                    "efficiency record'.")
    with rc2:
        rn = st.slider("Items", 10, 80, 40, key="radar_n")
    with rc3:
        rfresh_label = st.selectbox("Freshness",
                                    ["Past week", "Past 24 h",
                                     "Past month", "Any time"],
                                    key="radar_fresh",
                                    help="Only fetch items this recent - "
                                         "keeps the radar on the LATEST "
                                         "results instead of old stories.")
        rfresh = {"Past 24 h": "when:1d", "Past week": "when:7d",
                  "Past month": "when:30d", "Any time": ""}[rfresh_label]
    with rc4:
        rmodel = st.selectbox("Model", list(MODELS.keys()), index=0,
                              key="radar_model",
                              help="Haiku is ideal - cheap and accurate "
                                   "for short extraction.")

    wl = load_watch()
    can_call = (st.session_state.get("backend") == "max"
                or bool(api_key.strip()))

    b1, b2, b3 = st.columns([1.2, 1.6, 2.2])
    with b1:
        go_one = st.button("📡 Fetch & analyse", type="primary",
                           key="radar_go")
    with b2:
        go_all = st.button(
            f"🔁 Refresh all ({len(wl['queries'])} searches + "
            f"{len(wl['feeds'])} feeds + "
            f"{len(wl.get('patents', []))} patent watches)",
            key="radar_go_all",
            disabled=not (wl["queries"] or wl["feeds"]
                          or wl.get("patents")))
    with b3:
        if (rquery.strip() and rquery.strip() not in wl["queries"]
                and st.button("☆ Save this search as a watchlist",
                              key="radar_wl_add")):
            wl["queries"].append(rquery.strip())
            save_watch(wl)
            st.rerun()

    if wl["queries"]:
        wl_del = st.multiselect("Watchlists (select to remove)",
                                wl["queries"], key="radar_wl_del")
        if wl_del and st.button("Remove selected watchlists",
                                key="radar_wl_del_btn"):
            wl["queries"] = [q for q in wl["queries"] if q not in wl_del]
            save_watch(wl)
            st.rerun()

    if go_one or go_all:
        if not can_call:
            st.error("Extraction needs the API key (sidebar).")
        else:
            if st.session_state.get("backend") == "max":
                st.caption("Running on Claude Max — each news item is one "
                           "call, so large batches use up your Max window "
                           "faster. For 40+ items, Haiku in API-key mode "
                           "is the cheap option.")
            if go_all:
                jobs = ([("search", q) for q in wl["queries"]]
                        + [("feed", f) for f in wl["feeds"]]
                        + [("patent", q)
                           for q in wl.get("patents", [])])
            else:
                jobs = [("search",
                         rquery.strip() or "perovskite solar cell")]
            bar = st.progress(0.0, text="Fetching...")
            total_added = 0
            for qi, (kind, q) in enumerate(jobs):
                label = (q if kind == "search"
                         else q.split("://")[-1].split("/")[0])

                def _cb(i, n, title, _qi=qi, _q=label):
                    frac = (_qi + i / max(n, 1)) / len(jobs)
                    bar.progress(min(frac, 1.0),
                                 text=f"{_q}: {i}/{n} {title[:45]}")
                try:
                    if kind == "search":
                        added, size = run_fetch(q, rn, MODELS[rmodel],
                                                api_key.strip(), _cb,
                                                freshness=rfresh)
                    elif kind == "feed":
                        added, size = run_feed(q, rn, MODELS[rmodel],
                                               api_key.strip(), _cb)
                    else:
                        added, size = run_patents(q, 20,
                                                  MODELS[rmodel],
                                                  api_key.strip(), _cb)
                    total_added += added
                except Exception as e:
                    st.error(f"'{label}' failed: {e}")
            bar.empty()
            if total_added:
                st.success(f"Added {total_added} new item(s) across "
                           f"{len(jobs)} source(s). Store now "
                           f"{len(load_news())}. Switch to the Dashboard "
                           "to explore.")
            else:
                st.info(f"No new items (store has {len(load_news())}).")

    st.markdown("---")
    st.markdown("**🚨 Keyword alerts** — items matching these terms "
                "(title/company/summary) are flagged on the Dashboard: "
                "competitor names, 'certified record', a material "
                "system...")
    wl = load_watch()
    na1, na2 = st.columns([3, 1])
    with na1:
        new_alert = st.text_input("New alert term", key="alert_new",
                                  label_visibility="collapsed",
                                  placeholder="e.g. Renshine, certified "
                                              "record, CsPbI3")
    with na2:
        if (st.button("Add alert", key="alert_add",
                      use_container_width=True)
                and new_alert.strip()
                and new_alert.strip() not in wl["alerts"]):
            wl["alerts"].append(new_alert.strip())
            save_watch(wl)
            st.rerun()
    if wl["alerts"]:
        al_del = st.multiselect("Active alerts (select to remove)",
                                wl["alerts"], key="alert_del")
        if al_del and st.button("Remove selected alerts",
                                key="alert_del_btn"):
            wl["alerts"] = [a for a in wl["alerts"] if a not in al_del]
            save_watch(wl)
            st.rerun()

    st.markdown("---")
    st.markdown("**📶 Dedicated PV feeds** — specialist outlets pulled "
                "with every *Refresh all*, alongside Google News. Less "
                "noise, earlier catches. Any RSS URL works.")
    nf1, nf2 = st.columns([3, 1])
    with nf1:
        new_feed = st.text_input("New feed URL", key="feed_new",
                                 label_visibility="collapsed",
                                 placeholder="https://www.pv-magazine.com"
                                             "/feed/")
    with nf2:
        if (st.button("Add feed", key="feed_add",
                      use_container_width=True)
                and new_feed.strip().startswith("http")
                and new_feed.strip() not in wl["feeds"]):
            wl["feeds"].append(new_feed.strip())
            save_watch(wl)
            st.rerun()
    if wl["feeds"]:
        fd_del = st.multiselect("Active feeds (select to remove)",
                                wl["feeds"], key="feed_del")
        if fd_del and st.button("Remove selected feeds",
                                key="feed_del_btn"):
            wl["feeds"] = [f for f in wl["feeds"] if f not in fd_del]
            save_watch(wl)
            st.rerun()

    st.markdown("---")
    st.markdown("**🧾 Patent watch** — standing full-text searches on "
                "WIPO Patentscope. New filings land in the store with "
                "category *Patent*, so they surface on the Dashboard, "
                "in alerts, digests and investor reports like any other "
                "signal — early sight of where competitors are heading, "
                "before their papers appear.")
    pt1, pt2 = st.columns([3, 1])
    with pt1:
        new_pat = st.text_input("New patent search", key="pat_new",
                                label_visibility="collapsed",
                                placeholder="e.g. perovskite silicon "
                                            "tandem module")
    with pt2:
        if (st.button("Add patent watch", key="pat_add",
                      use_container_width=True)
                and new_pat.strip()
                and new_pat.strip() not in wl.get("patents", [])):
            wl.setdefault("patents", []).append(new_pat.strip())
            save_watch(wl)
            st.rerun()
    if wl.get("patents"):
        pt_del = st.multiselect("Patent watches (select to remove)",
                                wl["patents"], key="pat_del")
        if pt_del and st.button("Remove selected patent watches",
                                key="pat_del_btn"):
            wl["patents"] = [p for p in wl["patents"]
                             if p not in pt_del]
            save_watch(wl)
            st.rerun()
        if st.button("🧾 Check patents now", key="pat_go"):
            if not can_call:
                st.error("Extraction needs the API key (sidebar).")
            else:
                bar = st.progress(0.0, text="Searching Patentscope...")
                tot = 0
                for pi, q in enumerate(wl["patents"], start=1):
                    bar.progress(pi / len(wl["patents"]), text=q[:60])
                    try:
                        added, _ = run_patents(q, 20, MODELS[rmodel],
                                               api_key.strip())
                        tot += added
                    except Exception as e:
                        st.error(f"'{q}' failed: {e} — if this "
                                 "persists, tell Claude and the patent "
                                 "source can be switched.")
                bar.empty()
                if tot:
                    st.success(f"Added {tot} new filing(s).")
                else:
                    st.info("No new filings.")

# ------------------------------ TRENDS ------------------------------------
with tab_trend:
    if view is None:
        st.info(NO_DATA)
    else:
        tr = view.dropna(subset=["date_dt"]).copy()
        if tr.empty:
            st.caption("No dated items in this view yet.")
        else:
            tr["month"] = tr["date_dt"].dt.to_period("M").dt.to_timestamp()

            st.markdown("**Coverage over time** — analysed items per "
                        "month, by category.")
            mo = (tr.groupby(["month", "category"]).size()
                  .reset_index(name="items"))
            if alt is not None and not mo.empty:
                st.altair_chart(alt.Chart(mo).mark_bar().encode(
                    x=alt.X("month:T", title="Month"),
                    y=alt.Y("items:Q", title="Items", stack="zero"),
                    color=alt.Color("category:N", title="Category"),
                    tooltip=["month:T", "category:N", "items:Q"]),
                    use_container_width=True)

            tc1, tc2 = st.columns(2)
            with tc1:
                st.markdown("**Cell-type share of coverage**")
                ct = tr[tr["cell_type"].fillna("unspecified")
                        != "unspecified"]
                cts = (ct.groupby(["month", "cell_type"]).size()
                       .reset_index(name="items"))
                if alt is not None and not cts.empty:
                    st.altair_chart(alt.Chart(cts).mark_area().encode(
                        x=alt.X("month:T", title="Month"),
                        y=alt.Y("items:Q", stack="normalize",
                                title="Share"),
                        color=alt.Color("cell_type:N", title="Cell type"),
                        tooltip=["month:T", "cell_type:N", "items:Q"]),
                        use_container_width=True)
                else:
                    st.caption("No cell-type data yet.")
            with tc2:
                st.markdown("**Best reported efficiency per month**")
                ef = tr.dropna(subset=["efficiency_pct"])
                ef = ef[ef["efficiency_pct"].between(0.1, 50)]
                efm = (ef.groupby(["month", "cell_type"])
                       ["efficiency_pct"].max().reset_index())
                if alt is not None and not efm.empty:
                    st.altair_chart(alt.Chart(efm).mark_line(
                        point=True).encode(
                        x=alt.X("month:T", title="Month"),
                        y=alt.Y("efficiency_pct:Q", title="Best eff. (%)",
                                scale=alt.Scale(zero=False)),
                        color=alt.Color("cell_type:N", title="Cell type"),
                        tooltip=["month:T", "cell_type:N",
                                 "efficiency_pct:Q"]),
                        use_container_width=True)
                else:
                    st.caption("No efficiency values yet.")

            st.markdown("**Company momentum** — coverage in the last 90 "
                        "days vs the 90 days before. Risers are heating "
                        "up; fallers are going quiet.")
            now = tr["date_dt"].max()
            recent = tr[tr["date_dt"] >= now - _pd.Timedelta(days=90)]
            prior = tr[(tr["date_dt"] >= now - _pd.Timedelta(days=180))
                       & (tr["date_dt"] < now - _pd.Timedelta(days=90))]

            def _by_company(frame):
                c = frame.dropna(subset=["company"]).copy()
                c["company"] = c["company"].astype(str).str.strip()
                c = c[c["company"] != ""]
                return c.groupby("company").size()

            r_ct, p_ct = _by_company(recent), _by_company(prior)
            companies = sorted(set(r_ct.index) | set(p_ct.index))
            if not companies:
                st.caption("No company data yet.")
            else:
                mov = _pd.DataFrame({
                    "Company": companies,
                    "Last 90d": [int(r_ct.get(c, 0)) for c in companies],
                    "Prior 90d": [int(p_ct.get(c, 0)) for c in companies],
                })
                mov["Δ"] = mov["Last 90d"] - mov["Prior 90d"]
                mov["Trend"] = mov["Δ"].apply(
                    lambda d: "▲" if d > 0 else ("▼" if d < 0 else "—"))
                mov = mov.sort_values(["Δ", "Last 90d"],
                                      ascending=[False, False])
                st.dataframe(mov[["Company", "Trend", "Last 90d",
                                  "Prior 90d", "Δ"]].head(25),
                             use_container_width=True, height=340,
                             hide_index=True)

# ------------------------------ RECORDS -----------------------------------
with tab_rec:
    if view is None:
        st.info(NO_DATA)
    else:
        st.markdown("**Efficiency record progression** - the running best "
                    "reported value per cell type, from every analysed "
                    "item in the current view.")
        rec = view.dropna(subset=["efficiency_pct", "date_dt"]).copy()
        rec = rec[rec["efficiency_pct"].between(0.1, 50)]
        if rec.empty:
            st.caption("No dated efficiency values in this view yet.")
        else:
            rec = rec.sort_values("date_dt")
            rec["running_best"] = (rec.groupby("cell_type")
                                   ["efficiency_pct"].cummax())
            events = rec[rec["efficiency_pct"]
                         >= rec["running_best"]].copy()
            events = (events.groupby(["cell_type", "efficiency_pct"],
                                     as_index=False).first()
                      .sort_values("date_dt"))
            if alt is not None:
                line = alt.Chart(rec).mark_line(
                    interpolate="step-after", strokeWidth=2).encode(
                    x=alt.X("date_dt:T", title="Date"),
                    y=alt.Y("running_best:Q",
                            title="Best reported eff. (%)",
                            scale=alt.Scale(zero=False)),
                    color=alt.Color("cell_type:N", title="Cell type"))
                pts = alt.Chart(events).mark_point(
                    size=110, filled=True).encode(
                    x="date_dt:T", y="efficiency_pct:Q",
                    color="cell_type:N",
                    tooltip=["date:N", "cell_type:N", "efficiency_pct:Q",
                             "company:N", "title:N"])
                st.altair_chart(line + pts, use_container_width=True)
            st.markdown("**Record-setting items**")
            ev_show = (events.sort_values("date_dt", ascending=False)
                       [["date", "cell_type", "efficiency_pct", "company",
                         "title"]]
                       .rename(columns={"date": "Date",
                                        "cell_type": "Cell type",
                                        "efficiency_pct": "Eff %",
                                        "company": "Company",
                                        "title": "Headline"}))
            st.dataframe(ev_show, use_container_width=True, height=280)
            st.caption("Note: values come from news text, not "
                       "certification charts - verify a number before "
                       "quoting it.")

# ------------------------------ MAP ---------------------------------------
with tab_map:
    if view is None:
        st.info(NO_DATA)
    else:
        cc = view["country"].dropna().astype(str).str.strip()
        cc = cc[cc != ""]
        by_country = (cc.value_counts().rename_axis("country")
                      .reset_index(name="items"))
        if by_country.empty:
            st.caption("No country information in this view.")
        else:
            pts = []
            for _, r in by_country.iterrows():
                ll = country_latlon(r["country"])
                if ll:
                    pts.append({"country": r["country"],
                                "items": int(r["items"]),
                                "lat": ll[0], "lon": ll[1],
                                "radius": min(int(r["items"]) * 40000,
                                              500000)})
            if pts:
                mp = _pd.DataFrame(pts)
                try:
                    st.map(mp, latitude="lat", longitude="lon",
                           size="radius")
                except Exception:
                    try:
                        st.map(mp[["lat", "lon"]])
                    except Exception:
                        pass
            st.markdown("**Items by country**")
            if alt is not None:
                st.altair_chart(alt.Chart(by_country.head(20)).mark_bar()
                                .encode(x="items:Q",
                                        y=alt.Y("country:N", sort="-x"),
                                        tooltip=["country", "items"]),
                                use_container_width=True)
            else:
                st.bar_chart(by_country.set_index("country"))
            unmapped = [c for c in by_country["country"]
                        if not country_latlon(c)]
            if unmapped:
                st.caption("Not on map (no coordinates): "
                           + ", ".join(unmapped[:15]))

# ------------------------------ COMPANIES ---------------------------------
with tab_comp:
    if view is None:
        st.info(NO_DATA)
    else:
        comp = view.dropna(subset=["company"]).copy()
        comp["company"] = comp["company"].astype(str).str.strip()
        comp = comp[comp["company"] != ""]
        if comp.empty:
            st.caption("No companies detected in this view.")
        else:
            agg = (comp.groupby("company").agg(
                items=("title", "count"),
                best_eff=("efficiency_pct", "max"),
                funding=("funding_musd", "sum"),
                latest=("date", "max")).reset_index()
                .sort_values("items", ascending=False))
            st.dataframe(agg.rename(columns={
                "company": "Company", "items": "Items",
                "best_eff": "Best eff %", "funding": "Funding US$M",
                "latest": "Latest"}), use_container_width=True,
                height=340)
            if alt is not None:
                st.altair_chart(alt.Chart(agg.head(15)).mark_bar().encode(
                    x="items:Q", y=alt.Y("company:N", sort="-x"),
                    tooltip=["company", "items", "best_eff", "funding"]),
                    use_container_width=True)
            st.caption("For a written briefing on one company, use the "
                       "**📑 Reports** tab.")

# ------------------------------ REPORTS -----------------------------------
with tab_rep:
    if view is None:
        st.info(NO_DATA)
    else:
        st.markdown("**Report builder** - all written only from your "
                    "analysed items and saved to the answers folder "
                    "(auto-save toggle in the sidebar).")
        rt1, rt2 = st.columns([3, 1])
        with rt1:
            rtype = st.radio("Report type",
                             ["Digest (quick summary)",
                              "Company briefing",
                              "Detailed intelligence report",
                              "Investor report"],
                             horizontal=True, key="rep_type")
        with rt2:
            rep_model_label = st.selectbox(
                "Model for this report", list(MODELS.keys()), index=3,
                key="rep_model",
                help="Opus/Frontier for reports that matter (investor "
                     "material) - the stronger models synthesize and "
                     "hedge far better. Haiku/Sonnet for quick digests.")
        rep_model = MODELS[rep_model_label]

        can_call = (st.session_state.get("backend") == "max"
                    or bool(api_key.strip()))

        def _item_lines(frame, cap):
            lines = []
            for _, r in frame.sort_values("date_dt").tail(cap).iterrows():
                bits = [str(r.get("date") or ""),
                        str(r.get("category") or ""),
                        str(r.get("company") or ""),
                        str(r.get("country") or "")]
                if _pd.notna(r.get("efficiency_pct")):
                    bits.append(f"{r['efficiency_pct']}%")
                if _pd.notna(r.get("funding_musd")):
                    bits.append(f"US${r['funding_musd']}M")
                bits.append(str(r.get("title") or ""))
                lines.append(" | ".join(b for b in bits if b))
            return lines

        if rtype == "Digest (quick summary)":
            dg1, dg2 = st.columns([1, 3])
            with dg1:
                dg_days = st.selectbox("Period", [7, 14, 30],
                                       key="digest_days",
                                       format_func=lambda d:
                                       f"Last {d} days")
            recent_d = view.dropna(subset=["date_dt"])
            if not recent_d.empty:
                cutoff_d = (recent_d["date_dt"].max()
                            - _pd.Timedelta(days=dg_days))
                recent_d = recent_d[recent_d["date_dt"] >= cutoff_d]
            with dg2:
                st.caption(f"{len(recent_d)} item(s) in this period "
                           "(respects the global filter).")
            if st.button("📰 Build digest", type="primary",
                         key="digest_go", disabled=recent_d.empty):
                if not can_call:
                    st.error("Needs the API key (sidebar).")
                else:
                    period = f"last {dg_days} days"
                    umsg = (f"PERIOD: {period}\n\nITEMS:\n"
                            + "\n".join(_item_lines(recent_d, 120)))
                    with st.spinner(f"{rep_model_label} is writing..."):
                        try:
                            dg = call_claude(
                                api_key.strip(),
                                DIGEST_SYSTEM.replace("{period}", period),
                                umsg, rep_model, max_tokens=3000)
                        except Exception as e:
                            st.error(f"Claude error: {e}")
                            dg = None
                    if dg:
                        st.session_state["last_digest"] = dg
                        if do_autosave:
                            save_report(f"digest_{dg_days}d", dg)
            if st.session_state.get("last_digest"):
                st.markdown("---")
                st.markdown(st.session_state["last_digest"])
                st.download_button(
                    "⬇️ Digest as Markdown",
                    data=st.session_state["last_digest"].encode("utf-8"),
                    file_name="PVRadar_digest.md", mime="text/markdown")

        elif rtype == "Company briefing":
            names = sorted(view["company"].dropna().astype(str).str.strip()
                           .unique())
            names = [c for c in names if c]
            if not names:
                st.caption("No companies detected - fetch more or widen "
                           "filters.")
            else:
                pick = st.selectbox("Company", names,
                                    key="radar_company_pick")
                if st.button("📝 Generate briefing", type="primary",
                             key="radar_report"):
                    if not can_call:
                        st.error("Needs the API key (sidebar).")
                    else:
                        items = view[view["company"].astype(str).str.strip()
                                     == pick]
                        umsg = (f"ORGANISATION: {pick}\n\nNEWS ITEMS:\n"
                                + "\n".join(_item_lines(items, 60)))
                        with st.spinner(f"{rep_model_label} is briefing..."):
                            try:
                                rep = call_claude(api_key.strip(),
                                                  COMPANY_REPORT_SYSTEM,
                                                  umsg, rep_model)
                            except Exception as e:
                                st.error(f"Claude error: {e}")
                                rep = None
                        if rep:
                            st.session_state["last_brief"] = (pick, rep)
                            if do_autosave:
                                save_report(pick, rep)
                if st.session_state.get("last_brief"):
                    bname, btext = st.session_state["last_brief"]
                    st.markdown("---")
                    st.markdown(f"### Briefing: {bname}")
                    st.markdown(btext)
                    st.download_button(
                        "⬇️ Briefing as Markdown",
                        data=btext.encode("utf-8"),
                        file_name=f"PVRadar_{bname}.md",
                        mime="text/markdown")

        elif rtype == "Detailed intelligence report":
            rp1, rp2, rp3 = st.columns([2, 1, 1])
            with rp1:
                rp_title = st.text_input(
                    "Report title / focus", key="rep_title",
                    placeholder="e.g. Perovskite tandem landscape Q3 2026")
            with rp2:
                rp_days = st.selectbox("Period",
                                       ["30 days", "90 days", "12 months",
                                        "Everything"], index=1,
                                       key="rep_days")
            with rp3:
                rp_words = st.selectbox("Depth (words)",
                                        [800, 1500, 2500], index=1,
                                        key="rep_words")
            rp_focus = st.text_input(
                "Only items mentioning (optional, comma-separated)",
                key="rep_focus",
                placeholder="e.g. tandem, Oxford PV, encapsulation")
            rep_d = view.dropna(subset=["date_dt"])
            if rp_days != "Everything" and not rep_d.empty:
                nd = {"30 days": 30, "90 days": 90,
                      "12 months": 365}[rp_days]
                rep_d = rep_d[rep_d["date_dt"]
                              >= rep_d["date_dt"].max()
                              - _pd.Timedelta(days=nd)]
            if rp_focus.strip():
                terms = [t.strip().lower()
                         for t in rp_focus.split(",") if t.strip()]
                hay = (rep_d["title"].fillna("").astype(str) + " "
                       + rep_d["company"].fillna("").astype(str) + " "
                       + rep_d["summary"].fillna("").astype(str)
                       ).str.lower()
                mask = False
                for t in terms:
                    mask = mask | hay.str.contains(t, regex=False)
                rep_d = rep_d[mask]
            st.caption(f"{len(rep_d)} item(s) will feed this report "
                       "(global filter + period + focus).")
            if st.button("📑 Build detailed report", type="primary",
                         key="rep_go", disabled=rep_d.empty):
                if not can_call:
                    st.error("Needs the API key (sidebar).")
                else:
                    title = (rp_title.strip()
                             or f"PV intelligence report - {rp_days}")
                    sys_p = (REPORT_SYSTEM
                             .replace("{title}", title)
                             .replace("{words}", str(rp_words)))
                    umsg = ("ITEMS:\n"
                            + "\n".join(_item_lines(rep_d, 200)))
                    with st.spinner(f"{rep_model_label} is writing the "
                                    "report (this is the long one)..."):
                        try:
                            rp = call_claude(api_key.strip(), sys_p,
                                             umsg, rep_model,
                                             max_tokens=12000)
                        except Exception as e:
                            st.error(f"Claude error: {e}")
                            rp = None
                    if rp:
                        st.session_state["last_report"] = (title, rp)
                        if do_autosave:
                            save_report(title, rp)
            if st.session_state.get("last_report"):
                rtitle, rtext = st.session_state["last_report"]
                st.markdown("---")
                st.markdown(rtext)
                st.download_button(
                    "⬇️ Report as Markdown",
                    data=rtext.encode("utf-8"),
                    file_name=f"PVRadar_{rtitle[:40]}.md",
                    mime="text/markdown")

        else:  # Investor report
            st.caption("Investor-facing templates for startup material. "
                       "Built strictly from your collected items, with "
                       "figures traced to company + date and an explicit "
                       "data-limitations note — verify key numbers against "
                       "primary sources before they go in a deck.")
            iv1, iv2 = st.columns([2, 2])
            with iv1:
                iv_template = st.selectbox("Template",
                                           list(INVESTOR_TEMPLATES.keys()),
                                           key="iv_template")
            with iv2:
                iv_audience = st.text_input(
                    "Audience (shapes tone & focus)", key="iv_aud",
                    placeholder="e.g. seed-stage deep-tech VCs for a "
                                "perovskite module startup")
            iv3, iv4, iv5 = st.columns([2, 1, 1])
            with iv3:
                iv_title = st.text_input(
                    "Report title", key="iv_title",
                    placeholder="e.g. Perovskite PV competitive landscape "
                                "- Q3 2026")
            with iv4:
                iv_days = st.selectbox("Period",
                                       ["90 days", "6 months", "12 months",
                                        "Everything"], index=2,
                                       key="iv_days")
            with iv5:
                iv_words = st.selectbox("Depth (words)",
                                        [1000, 1800, 2500], index=1,
                                        key="iv_words")
            iv_focus = st.text_input(
                "Only items mentioning (optional, comma-separated)",
                key="iv_focus",
                placeholder="e.g. tandem, module, manufacturing")
            iv_d = view.dropna(subset=["date_dt"])
            if iv_days != "Everything" and not iv_d.empty:
                nd = {"90 days": 90, "6 months": 182,
                      "12 months": 365}[iv_days]
                iv_d = iv_d[iv_d["date_dt"]
                            >= iv_d["date_dt"].max()
                            - _pd.Timedelta(days=nd)]
            if iv_focus.strip():
                terms = [t.strip().lower()
                         for t in iv_focus.split(",") if t.strip()]
                hay = (iv_d["title"].fillna("").astype(str) + " "
                       + iv_d["company"].fillna("").astype(str) + " "
                       + iv_d["summary"].fillna("").astype(str)
                       ).str.lower()
                mask = False
                for t in terms:
                    mask = mask | hay.str.contains(t, regex=False)
                iv_d = iv_d[mask]
            st.caption(f"{len(iv_d)} item(s) will feed this report. For "
                       "investor material, more coverage = more credible: "
                       "run a broad fetch first if this number looks low.")
            if st.button("💼 Build investor report", type="primary",
                         key="iv_go", disabled=iv_d.empty):
                if not can_call:
                    st.error("Needs the API key (sidebar).")
                else:
                    title = (iv_title.strip()
                             or f"{iv_template} - perovskite PV "
                                f"({iv_days})")
                    audience = (iv_audience.strip()
                                or "startup founders preparing investor "
                                   "material")
                    sys_p = (INVESTOR_SYSTEM
                             .replace("{title}", title)
                             .replace("{audience}", audience)
                             .replace("{words}", str(iv_words))
                             .replace("{structure}",
                                      INVESTOR_TEMPLATES[iv_template]))
                    umsg = ("ITEMS:\n"
                            + "\n".join(_item_lines(iv_d, 250)))
                    with st.spinner(f"{rep_model_label} is writing the "
                                    "investor report (worth the wait)..."):
                        try:
                            iv_out = call_claude(api_key.strip(), sys_p,
                                                 umsg, rep_model,
                                                 max_tokens=12000)
                        except Exception as e:
                            st.error(f"Claude error: {e}")
                            iv_out = None
                    if iv_out:
                        st.session_state["last_investor"] = (title, iv_out)
                        if do_autosave:
                            save_report(title, iv_out)
            if st.session_state.get("last_investor"):
                ititle, itext = st.session_state["last_investor"]
                st.markdown("---")
                st.markdown(itext)
                st.download_button(
                    "⬇️ Report as Markdown",
                    data=itext.encode("utf-8"),
                    file_name=f"PVRadar_{ititle[:40]}.md",
                    mime="text/markdown")

# ------------------------------ BENCHMARK ---------------------------------
with tab_bench:
    st.markdown("**Startup benchmark** - your venture vs the competition, "
                "per product segment. You enter the facts; the radar "
                "enriches performance and funding from the news; Claude "
                "writes the positioning analysis. Empty cells stay "
                "'unknown' - nothing is ever invented.")
    st.caption("Stored locally in answers/pv_benchmark.json (covered by "
               "the nightly backup). Your startup's data leaves this "
               "machine only when you click the analysis button.")

    bseg = st.selectbox("Product segment", BENCH_SEGMENTS, key="bench_seg")
    bench = load_bench()
    b_rows = bench.get(bseg) or [_bench_row("My startup")]
    bdf = _pd.DataFrame(b_rows).reindex(columns=BENCH_COLS).fillna("")
    try:
        _bcfg = {
            "News: best eff %": st.column_config.TextColumn(
                disabled=True, help="Auto-filled by 'Enrich from radar' - "
                                    "best reported efficiency for this "
                                    "segment's cell type."),
            "News: funding US$M": st.column_config.TextColumn(
                disabled=True, help="Auto-filled by 'Enrich from radar' - "
                                    "summed reported funding."),
        }
    except Exception:
        _bcfg = None
    edited = st.data_editor(bdf, num_rows="dynamic",
                            use_container_width=True,
                            key=f"bench_ed_{bseg}", column_config=_bcfg)
    st.caption("Add a row per competitor (+ button below the table). "
               "Mark your own targets clearly, e.g. '26% target M24' vs "
               "'24.3% achieved'.")

    bb1, bb2, bb3 = st.columns(3)
    with bb1:
        if st.button("💾 Save table", key="bench_save",
                     use_container_width=True):
            recs = [r for r in edited.fillna("").astype(str)
                    .to_dict("records") if r["Company"].strip()]
            bench[bseg] = recs
            save_bench(bench)
            st.success(f"Saved {len(recs)} row(s) for {bseg}.")
    with bb2:
        if st.button("📡 Enrich from radar", key="bench_enrich",
                     use_container_width=True,
                     help="Fills the two News columns from the collected "
                          "items: best reported efficiency for this "
                          "segment's cell type, and summed funding. Your "
                          "own columns are never touched."):
            if view is None:
                st.info("No radar data yet - fetch news first.")
            else:
                target_ct = ("single junction"
                             if bseg == "Single junction" else "tandem")
                recs = [r for r in edited.fillna("").astype(str)
                        .to_dict("records") if r["Company"].strip()]
                for r in recs:
                    name = r["Company"].strip().lower()
                    m = (df["company"].fillna("").astype(str)
                         .str.lower().str.contains(name, regex=False))
                    if not m.any():
                        r["News: best eff %"] = "-"
                        r["News: funding US$M"] = "-"
                        continue
                    sub = df[m]
                    eff = sub[sub["cell_type"] == target_ct][
                        "efficiency_pct"].dropna()
                    r["News: best eff %"] = (f"{eff.max():.1f}"
                                             if not eff.empty else "-")
                    fund = sub["funding_musd"].dropna().sum()
                    r["News: funding US$M"] = (f"{fund:,.0f}"
                                               if fund else "-")
                bench[bseg] = recs
                save_bench(bench)
                st.success("News columns updated (4T/2T note: news "
                           "extraction only knows 'tandem', so both "
                           "tandem segments enrich from tandem items).")
                st.rerun()
    with bb3:
        bench_model_label = st.selectbox(
            "Analysis model", list(MODELS.keys()), index=3,
            key="bench_model", label_visibility="collapsed",
            help="Opus/Frontier recommended for investor material.")

    if st.button("🥊 Benchmark analysis (all segments)", type="primary",
                 key="bench_go"):
        can_call = (st.session_state.get("backend") == "max"
                    or bool(api_key.strip()))
        if not can_call:
            st.error("Needs the API key (sidebar).")
        else:
            bench = load_bench()
            seg_blocks, all_names = [], set()
            for seg in BENCH_SEGMENTS:
                recs = [r for r in bench.get(seg, [])
                        if str(r.get("Company", "")).strip()]
                if not recs:
                    continue
                lines = []
                for r in recs:
                    parts = [f"{c}: {str(r.get(c, '')).strip() or 'unknown'}"
                             for c in BENCH_COLS]
                    lines.append(" | ".join(parts))
                    all_names.add(str(r["Company"]).strip().lower())
                seg_blocks.append(f"### SEGMENT: {seg}\n" + "\n".join(lines))
            if not seg_blocks:
                st.warning("All segment tables are empty - fill and save "
                           "at least one first.")
            else:
                news_ctx = ""
                if view is not None:
                    nm = (df["company"].fillna("").astype(str).str.lower()
                          .apply(lambda c: any(n in c for n in all_names
                                               if n and n != "my startup")))
                    nsub = (df[nm].sort_values("date_dt", ascending=False)
                            .head(60))
                    if not nsub.empty:
                        nlines = []
                        for _, r in nsub.iterrows():
                            bits = [str(r.get("date") or ""),
                                    str(r.get("company") or "")]
                            if _pd.notna(r.get("efficiency_pct")):
                                bits.append(f"{r['efficiency_pct']}%")
                            if _pd.notna(r.get("funding_musd")):
                                bits.append(f"US${r['funding_musd']}M")
                            bits.append(str(r.get("title") or ""))
                            nlines.append(" | ".join(b for b in bits if b))
                        news_ctx = ("\n\nRECENT NEWS ITEMS FOR THESE "
                                    "COMPANIES:\n" + "\n".join(nlines))
                umsg = ("BENCHMARK TABLES (founder-maintained):\n\n"
                        + "\n\n".join(seg_blocks) + news_ctx)
                bm = MODELS[st.session_state.get("bench_model",
                                                 "Frontier (Fable 5.1)")]
                bml = st.session_state.get("bench_model", "Frontier (Fable 5.1)")
                with st.spinner(f"{bml} is writing the benchmark "
                                "analysis..."):
                    try:
                        out = call_claude(api_key.strip(),
                                          BENCHMARK_SYSTEM, umsg, bm,
                                          max_tokens=12000)
                    except Exception as e:
                        st.error(f"Claude error: {e}")
                        out = None
                if out:
                    st.session_state["last_bench"] = out
                    if do_autosave:
                        save_report("benchmark_analysis", out)
    if st.session_state.get("last_bench"):
        st.markdown("---")
        st.markdown(st.session_state["last_bench"])
        st.download_button(
            "⬇️ Analysis as Markdown",
            data=st.session_state["last_bench"].encode("utf-8"),
            file_name="PVRadar_benchmark_analysis.md",
            mime="text/markdown", key="bench_dl")

# ------------------------------ ARCHIVE -----------------------------------
with tab_arch:
    if view is None:
        st.info(NO_DATA)
    else:
        show_cols = {"date": "Date", "title": "Title",
                     "company": "Company", "country": "Country",
                     "efficiency_pct": "Eff %", "funding_musd": "US$M",
                     "category": "Category", "source": "Source"}
        tbl = (view.sort_values("date_dt", ascending=False)
               [[c for c in show_cols if c in view.columns]]
               .rename(columns=show_cols))
        st.dataframe(tbl, use_container_width=True, height=320)
        ax1, ax2, ax3 = st.columns(3)
        with ax1:
            st.download_button("⬇️ CSV (view)",
                               data=tbl.to_csv(index=False)
                               .encode("utf-8-sig"),
                               file_name="pv_radar.csv", mime="text/csv",
                               use_container_width=True)
        with ax2:
            ris = []
            for _, r in view.iterrows():
                yr = (str(r.get("date"))[:4] if r.get("date") else "")
                ris += ["TY  - NEWS", f"TI  - {r.get('title', '')}",
                        f"PY  - {yr}", f"DA  - {r.get('date', '')}",
                        f"PB  - {r.get('source', '')}",
                        f"UR  - {r.get('link', '')}",
                        f"KW  - {r.get('category', '')}", "ER  - ", ""]
            st.download_button("⬇️ Export .RIS (view)",
                               data="\n".join(ris).encode("utf-8"),
                               file_name="pv_radar.ris",
                               mime="application/x-research-info-systems",
                               use_container_width=True)
        with ax3:
            if st.button("🗑️ Clear entire store",
                         use_container_width=True):
                save_news([])
                st.rerun()
        del_pick = st.multiselect(
            "Delete specific items",
            options=[f"{r.get('date', '')} · "
                     f"{str(r.get('title', ''))[:70]}"
                     for _, r in view.iterrows()], key="radar_del")
        if del_pick and st.button("Delete selected", key="radar_del_btn"):
            del_titles = {p.split(" · ", 1)[-1] for p in del_pick}
            store = load_news()
            store = [r for r in store
                     if str(r.get("title", ""))[:70] not in del_titles]
            save_news(store)
            st.rerun()
